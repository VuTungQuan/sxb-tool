"""Tool SXB: quét kho file -> Excel, đối chiếu với 'Tong hop SXB 2006-2025.xlsx', phân loại,
kiểm tra nội dung file (năm trang cuối, lời nói đầu), đổi tên & lưu theo Năm/thang MM.

Chạy:  python -X utf8 sxb_tool.py            (giao diện)
       python -X utf8 sxb_tool.py --test     (tự kiểm tra)
"""
import csv, faulthandler, functools, glob, html, json, os, re, shutil, sqlite3, subprocess, sys, tempfile, threading, time, traceback, unicodedata, zipfile, datetime as dt
from collections import Counter, defaultdict

EXT_MAC_DINH = ".doc .docx .pdf .jpg .jpeg .png .tif .tiff .ai .psd .eps"
XL_TONG_HOP = "Tong hop SXB 2006-2025_chuan.xlsx"  # Excel tổng hợp mặc định (tìm cạnh exe hoặc thư mục cha)
XL_MA_BBT = "BBT_distinct.xlsx"  # mã BBT quy định: cột B = BBT (cột J Excel tổng hợp), cột F = mã dùng trong tên file
MA_BBT = {}  # {BBT nguyên gốc ('(trống)' = ô trống): mã quy định}, nạp lại mỗi lần đối chiếu
DUOI_BIA = {"ai", "jpg", "jpeg", "png", "tif", "tiff", "psd", "eps", "gif", "bmp"}  # ảnh / file thiết kế = bìa
WORD = {"doc", "docx"}
DOC_DUOC = {"pdf", "docx"}  # định dạng trích được text (.doc bỏ qua)
BIA = re.compile(r"(?:^| )(bia|to che|tr ?2 4)(?= |$)")  # file bìa / tờ che / trang 2-4 (đã chuan_hoa)
GIOI_HAN_O = 32000  # Excel chỉ chứa 32.767 ký tự/ô; dài hơn thì cắt theo dòng
KY_DAU, KY_CUOI = 25000, 8000  # .doc/.docx: ~8 trang đầu, ~2-3 trang cuối tính theo số ký tự (không bắt Word dàn trang)
GIAY_MOI_FILE = 0.25  # thời gian đọc 1 PDF/DOCX (đo trên ổ cục bộ), để ước tính trước khi quét
SO_LUONG = max(1, (os.cpu_count() or 2) // 2)  # số tiến trình đọc PDF/DOCX song song (~ số nhân thật: Xeon 16 nhân/32 luồng -> 16)
DUNG = threading.Event()  # nút Dừng: các vòng lặp dài kiểm tra cờ này sau mỗi file


class DaDung(Exception):
    pass


def _dung():
    if DUNG.is_set():
        raise DaDung("Đã dừng theo yêu cầu.")


def _gio(giay):
    """Số giây -> '45 giây' / '12 phút' / '3,5 giờ'."""
    if giay < 90:
        return f"{giay:.0f} giây"
    return f"{giay / 60:.0f} phút" if giay < 5400 else f"{giay / 3600:.1f} giờ".replace(".", ",")


def _ds_file(ds):
    """[(date_modified, path)] -> chuỗi 'Tất cả file tìm được': Word/PDF trước, ảnh sau (ô Excel bị cắt thì chỉ mất ảnh)."""
    ds = sorted(ds, key=lambda x: (not x[1].lower().endswith((".pdf", ".doc", ".docx")), x[1]))
    return "\n".join(f"{d}  {p}" for d, p in ds)


# ---------- chuẩn hoá tên ----------
def khong_dau(s):
    s = str(s).replace("đ", "d").replace("Đ", "D")
    s = unicodedata.normalize("NFD", s)
    return "".join(c for c in s if unicodedata.category(c) != "Mn")


def chuan_hoa(s):
    """'Lịch sử Chính phủ- T1' -> 'lich su chinh phu t1'"""
    return re.sub(r"[^a-z0-9]+", " ", khong_dau(s).lower()).strip()


def ma_ngay(v):
    """QĐXB -> ddmmyy hoặc mmyy. datetime, '01-2015', '1/2025', '19/1/2007', '11-2-18'."""
    if isinstance(v, (dt.datetime, dt.date)):
        return v.strftime("%d%m%y")
    nums = re.findall(r"\d+", str(v or ""))
    if len(nums) == 3:
        d, m, y = nums
        return f"{int(d):02d}{int(m):02d}{y[-2:]}"
    if len(nums) == 2:
        m, y = nums
        return f"{int(m):02d}{y[-2:]}"
    return ""


def ma_lan_xb(v):
    """None/1 -> M00 ; TB -> T01 ; số n>1 hoặc 'Lần 3'/'TB lần 2' -> T0n."""
    if v is None or str(v).strip() in ("", "1"):
        return "M00"
    s = str(v).strip().upper()
    n = re.findall(r"\d+", s)
    if n and int(n[0]) > 1:
        return f"T{int(n[0]):02d}"
    if s.startswith(("TB", "NB", "NỐI", "NOI", "LẦN", "LAN")):
        return "T01"
    return "M00"


DAI_TEN = 90  # độ dài tối đa phần tên sách: thư mục sách + tên file (đều = tên mới) giữ đường dẫn dưới 260 ký tự


CHU_THICH = re.compile(r" (?:sach (?:chuyen khao|tham khao)|xuat ban lan thu|tai ban|nguoi dich|bien dich|dich gia)\b.*$")


def loi_ten(ten):
    """Khoá để khớp tên sách: bỏ phần chú thích sau tên ('Sách chuyên khảo/tham khảo', 'Xuất bản lần thứ...', 'Tái bản...',
    'Người dịch / Biên dịch') vì tên file, tên thư mục thường không ghi. Giữ chú thích ngôn ngữ ('Tiếng Lào': sách khác)."""
    k = chuan_hoa(ten)
    g = CHU_THICH.sub("", k)
    return g if len(g.split()) >= 2 else k


def ten_file_moi(ten_sach, qdxb, bbt, lan_xb):
    """Tên mới 'ten_sach_0121_R_PL_T01'; phần tên sách dài quá DAI_TEN ký tự thì cắt ở ranh giới từ."""
    ten = "_".join(chuan_hoa(ten_sach).split())
    parts = [ten if len(ten) <= DAI_TEN else ten[:DAI_TEN + 1].rsplit("_", 1)[0]]
    if ma_ngay(qdxb):
        parts.append(ma_ngay(qdxb))
    ma = MA_BBT.get(str(bbt).strip() if bbt not in (None, "") else "(trống)")
    parts += ["R", ma or chuan_hoa(bbt or "").upper().replace(" ", "") or "XX", ma_lan_xb(lan_xb)]
    return "_".join(parts)


def doc_ma_bbt(path):
    """BBT_distinct.xlsx -> {BBT nguyên gốc: mã quy định (chỉ chữ in hoa không dấu + số)}; dòng chưa điền mã thì bỏ qua."""
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        rows = [r for r in wb.worksheets[0].iter_rows(min_row=2, values_only=True) if len(r) > 5 and r[1] is not None]
        # mã gõ dạng số (Excel bỏ số 0 đầu: 3 thay vì '03') -> 2 chữ số
        ma = lambda v: "" if v is None else f"{v:02d}" if isinstance(v, int) else re.sub(r"[^A-Z0-9]", "", chuan_hoa(str(v)).upper())
        return {k: v for k, v in ((str(r[1]).strip(), ma(r[5])) for r in rows) if v}
    finally:
        wb.close()


def thang_qdxb(v):
    """QĐXB 'mm/yyyy' hoặc 'dd/mm/yyyy' (hoặc datetime) -> 'MM'; thiếu -> '00'."""
    if isinstance(v, (dt.datetime, dt.date)):
        return v.strftime("%m")
    nums = re.findall(r"\d+", str(v or ""))
    return f"{int(nums[-2]):02d}" if len(nums) >= 2 else "00"


def duong_dan_dai(p):
    """Windows: thêm \\\\?\\ cho đường dẫn >240 ký tự."""
    if os.name == "nt" and len(p) > 240 and not p.startswith("\\\\?\\"):
        return "\\\\?\\" + os.path.abspath(p)
    return p


# ---------- quét kho ----------
COT_KHO = ["STT", "Tên file", "Định dạng", "Ngày sửa", "Năm", "Kích thước (KB)", "Đường dẫn"]
RONG_KHO = [6, 70, 10, 17, 7, 14, 100]
MAX_DONG = 1_048_575  # ponytail: Excel giới hạn 1.048.576 dòng/sheet -> quá thì mở sheet "Kho 2", "Kho 3"...


def _doc_nhat_ky(nhat_ky):
    """Đọc nhật ký quét: lặp (loại, các cột) của các thư mục đã xong; khối cuối chưa có dòng 'D' (đang quét dở thì tool tắt)
    bị bỏ và cắt khỏi file. -> danh sách dòng hợp lệ."""
    hop_le, cho, cuoi = [], [], 0
    with open(nhat_ky, "rb") as f:
        while True:
            dong = f.readline()
            if not dong:
                break
            if not dong.endswith(b"\n"):  # dòng ghi dở
                break
            try:
                row = next(csv.reader([dong.decode("utf-8")]))
            except (UnicodeDecodeError, csv.Error, StopIteration):
                break
            cho.append(row)
            if row and row[0] in ("#", "D"):
                hop_le += cho; cho = []; cuoi = f.tell()
    os.truncate(nhat_ky, cuoi)
    return hop_le


def quet(thu_muc, exts, xlsx_path, progress=None, bao_cao=None, nhat_ky=None):
    """Quét kho -> file .xlsx có định dạng (header đậm, cố định dòng 1, độ rộng cột, wrap text).
    Ghi dần vào nhật ký quét (CSV; mỗi thư mục quét xong là ghi xuống đĩa). Tool tắt đột ngột hoặc bấm Dừng giữa chừng:
    gọi lại với cùng nhat_ky thì bỏ qua các thư mục đã quét, làm tiếp. Xong thì chuyển sang .xlsx và xoá nhật ký.
    bao_cao (dict) nhận: 'loi' = [thư mục/file không đọc được], 'dinh_dang' = Counter, 'tiep' = số file lấy lại từ lần dở."""
    bao_cao = {} if bao_cao is None else bao_cao
    loi, dem = bao_cao.setdefault("loi", []), bao_cao.setdefault("dinh_dang", Counter())
    bo = lambda x: str(x or "").replace("\\\\?\\", "", 1)
    nhat_ky = nhat_ky or xlsx_path + ".dang_quet.csv"
    exts = {e if e.startswith(".") else "." + e for e in exts}
    goc = thu_muc
    if os.name == "nt" and not goc.startswith("\\\\?\\"):
        goc = "\\\\?\\" + os.path.abspath(goc)  # đường dẫn dài >260 ký tự
    dau = ["#", os.path.abspath(thu_muc), " ".join(sorted(exts))]
    da_xong, n = set(), 0
    if os.path.exists(nhat_ky):
        cu = _doc_nhat_ky(nhat_ky)
        if cu and cu[0] != dau:
            raise ValueError(f"Nhật ký quét dở là của thư mục/định dạng khác: {cu[0][1:]}")
        for row in cu[1:]:
            if row[0] == "F" and len(row) == 6:
                n += 1; dem[row[2].upper()] += 1
            elif row[0] == "D" and len(row) == 2:
                da_xong.add(row[1])
            elif row[0] == "L" and len(row) == 2:
                loi.append(row[1])
        if not cu:
            os.remove(nhat_ky)
    bao_cao["tiep"] = n
    moi = not os.path.exists(nhat_ky)
    with open(nhat_ky, "a", encoding="utf-8", newline="") as f:
        w = csv.writer(f)
        if moi:
            w.writerow(dau); f.flush()
        def loi_walk(e):  # thư mục không đọc được (không có quyền, đường dẫn hỏng...)
            loi.append(f"{bo(e.filename)}: {e.strerror}"); w.writerow(["L", loi[-1]])
        for root, _dirs, files in os.walk(goc, onerror=loi_walk):
            if bo(root) in da_xong:
                continue
            khoi = []  # cả thư mục ghi một lần kèm dòng 'D': tắt giữa chừng thì lần sau quét lại thư mục này từ đầu
            for fn in files:
                _dung()  # dừng: phần đã quét giữ trong nhật ký, chạy lại sẽ quét tiếp
                stem, ext = os.path.splitext(fn)
                ext = ext.lower()
                if exts and ext not in exts:
                    continue
                p = os.path.join(root, fn)
                try:
                    st = os.stat(p)
                except OSError as e:
                    khoi.append(["L", f"{bo(p)}: {e.strerror}"]); loi.append(khoi[-1][1])
                    continue
                dem[ext.lstrip(".").upper()] += 1; n += 1
                khoi.append(["F", stem, ext.lstrip(".").upper(), str(st.st_mtime), str(st.st_size), bo(p)])
                if progress and n % 1000 == 0:
                    progress(n, root)
            w.writerows(khoi + [["D", bo(root)]]); f.flush()

    # nhật ký -> Excel
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Alignment, Font
    wb = Workbook(write_only=True)
    wrap = Alignment(wrap_text=True, vertical="top")

    def o(ws, v, **kw):
        c = WriteOnlyCell(ws, value=v)
        for k, x in kw.items(): setattr(c, k, x)
        return c

    def sheet_moi():
        ws = wb.create_sheet(f"Kho {len(wb.sheetnames) + 1}" if wb.sheetnames else "Kho")
        for cot, r in zip("ABCDEFG", RONG_KHO):
            ws.column_dimensions[cot].width = r
        ws.freeze_panes = "A2"
        ws.append([o(ws, h, font=Font(bold=True), alignment=Alignment(horizontal="center")) for h in COT_KHO])
        return ws

    ws, n = sheet_moi(), 0
    for row in _doc_nhat_ky(nhat_ky):
        if row[0] != "F" or len(row) != 6:
            continue
        _, stem, ext, mtime, size, path = row
        d = dt.datetime.fromtimestamp(float(mtime))
        n += 1
        if n % MAX_DONG == 1 and n > 1:
            ws = sheet_moi()
        ws.append([n, o(ws, stem, alignment=wrap), ext,
                   o(ws, d.replace(second=0, microsecond=0), number_format="dd/mm/yyyy hh:mm"),
                   d.year, round(int(size) / 1024), o(ws, path, alignment=wrap)])
    wb.save(xlsx_path)
    os.remove(nhat_ky)
    return n


# ---------- đọc excel ----------
def _khoa_header(c):
    h = chuan_hoa(c)
    if "ten sach" in h: return "ten"
    if h.startswith("lan"): return "lan"
    if h in ("bbt", "ban bt", "ban bien tap"): return "bbt"
    if "qdxb" in h or "thoi gian" in h: return "qdxb"
    if h == "tt": return "tt"
    return None


def doc_excel(xlsx):
    """-> list dict {nam, dong, tt, ten, lan, bbt, qdxb}. Cột dò theo tên header (vị trí đổi theo năm)."""
    import openpyxl
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    sach = []
    for sh in wb.sheetnames:
        if not sh.strip().isdigit():
            continue
        nam, cot = int(sh), None
        for i, row in enumerate(wb[sh].iter_rows(values_only=True), 1):
            if cot is None:
                if any(c and "ten sach" in chuan_hoa(c) for c in row):
                    cot = {}
                    for idx, c in enumerate(row):
                        k = c and _khoa_header(c)
                        if k and k not in cot:
                            cot[k] = idx
                continue
            g = lambda k: row[cot[k]] if k in cot and cot[k] < len(row) else None
            ten = g("ten")
            if not ten or not str(ten).strip() or not chuan_hoa(ten):
                continue
            sach.append(dict(nam=nam, dong=i, tt=g("tt"), ten=str(ten).strip(),
                             lan=g("lan"), bbt=g("bbt"), qdxb=g("qdxb")))
    wb.close()
    if not sach:
        raise ValueError(f"'{os.path.basename(xlsx)}' không phải file Excel tổng hợp: cần các sheet năm (2006, 2007...) "
                         "có cột 'Tên sách'. Có thể đang chọn nhầm file kho vào ô này.")
    return sach


# ---------- đối chiếu ----------
def doc_kho(path):
    """Đọc file kho: .xlsx do quet() xuất, hoặc .csv cũ (utf-8 phẩy / utf-16 tab, header tiếng Việt hoặc máy)
    -> dict nội bộ: ten_file, dinh_dang (thường), date_modified (yyyy-mm-dd HH:MM:SS, để sort), path, key, nam."""
    rows = []
    if path.lower().endswith(".xlsx"):
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True)
        for ws in wb.worksheets:
            it = ws.iter_rows(values_only=True)
            if list(next(it, None) or ())[:len(COT_KHO)] != COT_KHO:
                wb.close()
                raise ValueError(f"'{os.path.basename(path)}' không phải file kho do tab 1 xuất (dòng đầu phải là: "
                                 f"{', '.join(COT_KHO)}). Có thể đang chọn nhầm file Excel tổng hợp vào ô này.")
            for r in it:
                _dung()
                if r and r[1] is not None:
                    rows.append(_dong_kho(str(r[1]), str(r[2]), r[3].strftime("%Y-%m-%d %H:%M:%S"), str(r[6])))
        wb.close()
    else:
        with open(path, "rb") as f:
            utf16 = f.read(2) == b"\xff\xfe"
        with open(path, encoding="utf-16" if utf16 else "utf-8-sig", newline="") as f:
            if not utf16 and not f.readline().startswith("sep="):
                f.seek(0)
            for r in csv.DictReader(f, delimiter="\t" if utf16 else ","):
                if "Tên file" in r:
                    d = dt.datetime.strptime(r["Ngày sửa"], "%d/%m/%Y %H:%M")
                    rows.append(_dong_kho(r["Tên file"], r["Định dạng"], d.strftime("%Y-%m-%d %H:%M:%S"), r["Đường dẫn"]))
                else:
                    rows.append(_dong_kho(r["ten_file"], r["dinh_dang"], r["date_modified"], r["path"]))
    return rows


def _dong_kho(ten_file, dinh_dang, date_modified, path):
    """Một dòng kho dạng gọn (~0,6 KB): chỉ giữ tên đã chuẩn hoá. ponytail: 4 triệu file vẫn cần vài GB RAM."""
    return dict(key=chuan_hoa(ten_file), dinh_dang=sys.intern(dinh_dang.lower()), date_modified=date_modified,
                path=path, nam=int(date_modified[:4]))


def key_thu_muc(thu_muc):
    """Tên thư mục này và các thư mục cha (chuẩn hoá, bỏ ổ đĩa và hậu tố bản sao "(1)")."""
    return [chuan_hoa(re.sub(r"\s*\(\d+\)\s*$", "", x)) for x in re.split(r"[\\/]", thu_muc)
            if x and not x.endswith(":") and x != "?"]


def la_bia(path, key_sach=""):
    """File bìa / tờ che / trang 2-4 — trừ khi chính tên sách có từ đó. Không xét tên thư mục: kho có thư mục 'Bia'
    chứa cả file ruột (…R06M00.pdf 19 trang)."""
    return bool(BIA.search(chuan_hoa(os.path.splitext(os.path.basename(path))[0]))) and not BIA.search(key_sach)


@functools.lru_cache(maxsize=200000)
def so_trang_pdf(path):
    """Số trang PDF (None nếu không mở được)."""
    try:
        import fitz
        with fitz.open(duong_dan_dai(path)) as d:
            return len(d)
    except Exception:
        return None


@functools.lru_cache(maxsize=200000)
def pdf_mong(path):
    """PDF bìa, không phải ruột sách: 1-2 trang (bìa 1, bìa 4, tờ che, trang 2-4), hoặc ≤ 8 trang mà trang đầu
    nằm ngang (trải bìa 1 + gáy + bìa 4, các trang sau là tai gập)."""
    if not path.lower().endswith(".pdf"):
        return False
    try:
        import fitz
        with fitz.open(duong_dan_dai(path)) as d:
            n, r = len(d), (d[0].rect if len(d) else None)
    except Exception:
        return False
    return 0 < n <= 2 or (n <= 8 and r.width > 1.3 * r.height)


def chon_file(ds, key_sach=""):
    """ds: [(date_modified, path)] -> (word, pdf, số file bìa bỏ qua). Mới nhất; cùng ngày sửa thì PDF nhiều trang hơn,
    rồi file lớn hơn. Bỏ file bìa (tên 'Bìa / Tờ che / Tr 2-4') và PDF 1-2 trang (chỉ đếm trang PDF của nhóm đang xét)."""
    def kt(p):
        try: return os.path.getsize(duong_dan_dai(p))
        except OSError: return 0
    chinh = [x for x in ds if not la_bia(x[1], key_sach)]
    w = [x for x in chinh if x[1].rsplit(".", 1)[-1].lower() in WORD]
    p = [x for x in chinh if x[1].lower().endswith(".pdf")]
    def best(xs):
        for ngay in sorted({x[0] for x in xs}, reverse=True):
            cung = [x for x in xs if x[0] == ngay and not pdf_mong(x[1])]
            if cung:
                trang = lambda x: (so_trang_pdf(x[1]) or 0) if x[1].lower().endswith(".pdf") else 0
                return max(cung, key=lambda x: (trang(x), kt(x[1])))[1]
        return ""
    return best(w), best(p), len(ds) - len(chinh)


def ds_bia(r):
    """File bìa của một dòng KHỚP: trong thư mục của file ruột đã chọn (sách chỉ có bìa: thư mục của file mới nhất),
    là ảnh / file thiết kế (.ai, .jpg, .png...), PDF bìa (pdf_mong) hoặc file bìa theo tên (la_bia, kể cả ở thư mục con 'Bia')."""
    ds = _dong_file(r["cac_file"])
    tm = {os.path.dirname(x) for x in (r["file_word"], r["file_pdf"]) if x} or {os.path.dirname(max(ds)[1])} if ds else set()
    chon, key, kq = {r["file_word"], r["file_pdf"]}, loi_ten(r["ten"]), []
    for _, p in ds:
        duoi, d = os.path.splitext(p)[1].lower().lstrip("."), os.path.dirname(p)
        if p not in chon and ((d in tm and (duoi in DUOI_BIA or (duoi == "pdf" and pdf_mong(p))))
                              or (any(d == t or d.startswith(t + os.sep) for t in tm) and la_bia(p, key))):
            kq.append(p)
    return sorted(set(kq))


def ten_bia(ten_moi):
    """Tên file bìa: đổi _R_ (Ruột) thành _B_ (Bìa) - 'ten_sach_0121_R_PL_T01' -> 'ten_sach_0121_B_PL_T01'."""
    return re.sub(r"_R_(?=[A-Z0-9]+_[MT]\d\d$)", "_B_", ten_moi, count=1)


def doi_chieu(sach, files):
    """Khớp tên sách với tên file HOẶC tên thư mục chứa file (kho xếp mỗi sách một thư mục)."""
    theo_key, ten_sach = set(), {loi_ten(s["ten"]) for s in sach}
    theo_prefix = defaultdict(list)  # ponytail: so khớp "chứa" chỉ trong nhóm cùng 3 từ đầu, tránh O(sách×file)
    theo_tm = defaultdict(list)
    for f in files:
        theo_key.add(f["key"])
        theo_prefix[" ".join(f["key"].split()[:3])].append((f["key"], False, f))
        theo_tm[os.path.dirname(f["path"])].append(f)
    for tm in theo_tm:  # tên thư mục: mỗi thư mục một mục chỉ mục cho nó và từng thư mục cha, không lặp theo file
        for k in set(key_thu_muc(tm)):
            theo_key.add(k)
            theo_prefix[" ".join(k.split()[:3])].append((k, True, tm))

    kq = []
    for s in sach:
        key = loi_ten(s["ten"])
        tim, rut_gon = {}, False
        for k, la_tm, x in theo_prefix.get(" ".join(key.split()[:3]), ()):
            khop = f" {key} " in f" {k} "  # gồm cả khớp chính xác
            gon = not khop and la_tm and len(k.split()) >= 3 and key.startswith(k + " ") and k not in ten_sach  # thư mục tên rút gọn
            if khop or gon:
                tim.update((f["path"], f) for f in (theo_tm[x] if la_tm else (x,)))
                rut_gon |= gon
        tim = list(tim.values())
        kieu = "chính xác" if key in theo_key else "thư mục rút gọn" if rut_gon else "chứa tên" if tim else ""
        khop_nam = [f for f in tim if f["nam"] == s["nam"]]

        r = dict(s, ten_moi=ten_file_moi(s["ten"], s["qdxb"], s["bbt"], s["lan"]),
                 kieu_khop=kieu, so_file=len(tim), file_word="", file_pdf="", ghi_chu="",
                 cac_file=_ds_file((f["date_modified"], f["path"]) for f in tim))
        if not tim:
            r["tinh_trang"] = "KHÔNG TÌM THẤY FILE"
        elif not khop_nam:
            r["tinh_trang"] = "LỆCH NĂM - cần mở file kiểm tra"
        else:
            r["tinh_trang"] = "KHỚP"
            r["file_word"], r["file_pdf"], bia = chon_file([(f["date_modified"], f["path"]) for f in khop_nam], key)
            r["ghi_chu"] = "; ".join(x for x in (
                f"{len(tim) - len(khop_nam)} file khác năm bị bỏ qua" if len(khop_nam) < len(tim) else "",
                f"bỏ {bia} file bìa/tờ che" if bia else "",
                "chỉ có file bìa" if not (r["file_word"] or r["file_pdf"]) else "") if x)
        kq.append(r)
    return kq


COT_KQ = [("nam", "Năm"), ("dong", "Dòng Excel"), ("tt", "TT"), ("ten", "Tên sách"), ("lan", "Lần XB"),
          ("bbt", "BBT"), ("qdxb", "QĐXB"), ("tinh_trang", "Tình trạng"), ("kieu_khop", "Kiểu khớp"),
          ("so_file", "Số file"), ("file_word", "File Word chọn"), ("file_pdf", "File PDF chọn"),
          ("ten_moi", "Tên file mới"), ("ghi_chu", "Ghi chú"), ("cac_file", "Tất cả file tìm được"),
          ("dich", "Đường dẫn đích")]


COT_TM = [("thu_muc", "Thư mục"), ("ten_tm", "Tên thư mục"), ("so_file", "Số file"), ("loai", "Loại file"),
          ("tinh_trang", "Tình trạng"), ("nam_cuoi", "Năm XB đọc được"), ("sach_khop", "Sách khớp"),
          ("sach_khac", "Sách trùng tên chưa khớp (năm - dòng Excel)"), ("ghi_chu", "Ghi chú")]
TM_KHOP, TM_CHUA, TM_KHONG = "KHỚP", "CHƯA KHỚP - chỉ trùng tên", "KHÔNG KHỚP SÁCH NÀO"
TM_TRUNG = "BẢN TRÙNG - sách đã khớp ở thư mục khác"


def doi_chieu_thu_muc(kq, files, cache=None):
    """Mỗi thư mục trong kho -> có file được chọn cho sách nào (KHỚP), là bản trùng của sách đã khớp nơi khác,
    chỉ trùng tên sách, hay không khớp sách nào. Năm trang cuối lấy từ cache tab 3 (không mở file ở đây)."""
    cache = {} if cache is None else cache
    theo_tm, khop, khac = defaultdict(list), defaultdict(list), defaultdict(list)
    for f in files:
        theo_tm[os.path.dirname(f["path"])].append(f)
    for r in kq:
        nhan = f"{r['nam']} - dòng {r['dong']} - {r['ten']}"
        chon = {os.path.dirname(x) for x in (r["file_word"], r["file_pdf"]) if x}
        for tm in chon:
            khop[tm].append(nhan)
        for tm in {os.path.dirname(x) for _, x in _dong_file(r["cac_file"])} - chon:
            khac[tm].append((nhan, r))
    kq_tm = []
    for tm, fs in sorted(theo_tm.items()):
        doc = [f["path"] for f in fs if f["dinh_dang"] in ("pdf", "doc", "docx") and not la_bia(f["path"])]
        ts = [cache.get(_khoa_cache(x), "chưa đọc") for x in doc]
        cuoi = sorted({y for t in ts if isinstance(t, dict) for y in t["nam"]})
        dau = sorted({y for t in ts if isinstance(t, dict) for y in t.get("dau", ())})
        nam = sorted(set(cuoi) | set(dau))
        doc_duoc, chua = sum(isinstance(t, dict) for t in ts), ts.count("chưa đọc")
        n_ocr = sum(isinstance(t, dict) and bool(t.get("ocr")) and not pdf_mong(x) for x, t in zip(doc, ts))  # PDF bìa không cần OCR
        ghi = (["không có file Word/PDF (chỉ ảnh, bìa, zip...)"] if not doc else
               [f"mở được {doc_duoc}/{len(doc)} file Word/PDF"]
               + (["còn lại không mở được (file hỏng, có mật khẩu, hoặc máy không có Word)"] if doc_duoc + chua < len(doc) else [])
               + ([f"{chua} file chưa đọc - chạy tab 3"] if chua else [])
               + ([f"{n_ocr} file PDF cần OCR (không có lớp chữ) - xem sheet CAN_OCR"] if n_ocr else [])
               + (["không thấy năm ở trang cuối lẫn 8 trang đầu (bản chính nằm trong .zip?)"] if doc_duoc and not nam and not n_ocr else []))
        trung = any(r["tinh_trang"] == "KHỚP" and r["nam"] in nam for _, r in khac[tm])
        kq_tm.append(dict(thu_muc=tm, ten_tm=os.path.basename(tm), so_file=len(fs),
                          loai=", ".join(f"{k.upper()} {n}" for k, n in sorted(Counter(f["dinh_dang"] for f in fs).items())),
                          tinh_trang=TM_KHOP if khop[tm] else TM_TRUNG if trung else TM_CHUA if khac[tm] else TM_KHONG,
                          nam_cuoi="; ".join(f"{n}: {', '.join(map(str, ys))}" for n, ys in (("trang cuối", cuoi), ("8 trang đầu", dau)) if ys),
                          sach_khop="\n".join(khop[tm]),
                          sach_khac="\n".join(n for n, _ in khac[tm]), ghi_chu="; ".join(ghi)))
    return kq_tm


COT_OCR = [("path", "Đường dẫn"), ("ten_tm", "Thư mục"), ("ten_file", "Tên file"), ("trang", "Số trang"),
           ("ly_do", "Lý do"), ("uu_tien", "Ưu tiên"), ("sach", "Sách liên quan (năm - dòng Excel - tên sách)")]
LY_DO_OCR = {"scan": "PDF ảnh scan, không có lớp chữ", "hinh_ve": "PDF chữ dạng hình vẽ (outline), không có lớp chữ"}
OCR_CAO, OCR_THAP = "CAO - thư mục chưa khớp sách nào", "thấp - thư mục đã khớp sách nhờ file khác"


def ds_can_ocr(kq, files, cache=None):
    """PDF đã đọc ở tab 3 mà gần như không có chữ -> cần OCR mới đọc được năm / lời nói đầu. Bỏ file bìa.
    -> danh sách cho sheet CAN_OCR (không mở file ở đây, lấy từ cache); file ở thư mục chưa khớp sách nào xếp trước."""
    cache = {} if cache is None else cache
    sach, da_khop = defaultdict(list), set()
    for r in kq:
        for _, p in _dong_file(r["cac_file"]):
            sach[p].append(f"{r['nam']} - dòng {r['dong']} - {r['ten']}")
        if r["tinh_trang"] == "KHỚP":
            da_khop |= {os.path.dirname(x) for x in (r["file_word"], r["file_pdf"]) if x}
    ds = []
    for f in files:
        if f["dinh_dang"] != "pdf" or la_bia(f["path"]):
            continue
        t = cache.get(_khoa_cache(f["path"]))
        if isinstance(t, dict) and t.get("ocr") and not pdf_mong(f["path"]):  # PDF bìa không cần OCR
            ds.append(dict(path=f["path"], ten_tm=os.path.basename(os.path.dirname(f["path"])), ten_file=os.path.basename(f["path"]),
                           trang=t.get("trang"), ly_do=LY_DO_OCR.get(t["ocr"], t["ocr"]),
                           uu_tien=OCR_THAP if os.path.dirname(f["path"]) in da_khop else OCR_CAO, sach="\n".join(sach[f["path"]])))
    return sorted(ds, key=lambda x: (x["uu_tien"] != OCR_CAO, x["path"]))


def _doc_sheet(xlsx, ten, cot):
    import openpyxl
    wb = openpyxl.load_workbook(xlsx, read_only=True)
    ten2k, rows = {h: k for k, h in cot}, []
    if ten in wb.sheetnames:
        it = wb[ten].iter_rows(values_only=True)
        keys = [ten2k.get(h) for h in next(it, ())]
        rows = [{k: ("" if v is None else v) for k, v in zip(keys, row) if k} for row in it]
    wb.close()
    return rows


def doc_thu_muc(xlsx):
    """Đọc lại sheet THU_MUC_KHO của file doi_chieu (rỗng nếu file cũ không có)."""
    return _doc_sheet(xlsx, "THU_MUC_KHO", COT_TM)


def doc_can_ocr(xlsx):
    """Đọc lại sheet CAN_OCR của file doi_chieu (rỗng nếu không có)."""
    return _doc_sheet(xlsx, "CAN_OCR", COT_OCR)


def ghi_ket_qua(kq, xlsx_goc, thu_muc_ra, thu_muc=(), can_ocr=()):
    import openpyxl
    os.makedirs(thu_muc_ra, exist_ok=True)
    tem = dt.datetime.now().strftime("%Y%m%d_%H%M")

    wb = openpyxl.Workbook()
    for ten, loc in (("KHOP", lambda r: r["tinh_trang"] == "KHỚP"),
                     ("CAN_XU_LY", lambda r: r["tinh_trang"] != "KHỚP")):
        ws = wb.create_sheet(ten)
        ws.append([h for _, h in COT_KQ])
        for r in kq:
            if loc(r):
                ws.append([_cell(r.get(k, "")) for k, _ in COT_KQ])
    for ten, rows, cot in (("THU_MUC_KHO", thu_muc, COT_TM), ("CAN_OCR", can_ocr, COT_OCR)):
        if rows:
            ws = wb.create_sheet(ten)
            ws.append([h for _, h in cot])
            for r in rows:
                ws.append([_cell(r.get(k, "")) for k, _ in cot])
            ws.auto_filter.ref = ws.dimensions
            ws.freeze_panes = "A2"
    del wb["Sheet"]
    ra1 = os.path.join(thu_muc_ra, f"doi_chieu_{tem}.xlsx")
    wb.save(ra1)

    # đánh dấu tình trạng vào bản sao file Excel tổng hợp; dòng KHỚP tô vàng cả dòng
    from openpyxl.styles import PatternFill
    vang = PatternFill("solid", fgColor="FFFF00")
    wb = openpyxl.load_workbook(xlsx_goc)
    theo_nam = defaultdict(list)
    for r in kq:
        theo_nam[r["nam"]].append(r)
    for nam, rs in theo_nam.items():
        ws = wb[str(nam)]
        c, hd = ws.max_column + 1, min(r["dong"] for r in rs) - 1
        for i, h in enumerate(("Tình trạng", "File chọn", "Tên file mới", "Đường dẫn đích")):
            ws.cell(hd, c + i, h)
        for r in rs:
            ws.cell(r["dong"], c, r["tinh_trang"])
            ws.cell(r["dong"], c + 1, _cell("\n".join(x for x in (r["file_word"], r["file_pdf"]) if x)))
            ws.cell(r["dong"], c + 2, r["ten_moi"])
            ws.cell(r["dong"], c + 3, _cell(r.get("dich", "")))
            if r["tinh_trang"] == "KHỚP":
                for cot in range(1, c + 4):
                    ws.cell(r["dong"], cot).fill = vang
    ra2 = os.path.join(thu_muc_ra, f"{os.path.splitext(os.path.basename(xlsx_goc))[0]}_danh_dau_{tem}.xlsx")
    wb.save(ra2)
    return ra1, ra2


def _cell(v):
    if isinstance(v, (dt.datetime, dt.date)):
        return v.strftime("%d/%m/%Y")
    if isinstance(v, str) and len(v) > GIOI_HAN_O:  # Excel báo lỗi file nếu ô dài hơn 32.767 ký tự
        giu = v[:GIOI_HAN_O].split("\n")[:-1]
        bo = v.count("\n") + 1 - len(giu)
        return "\n".join(giu or [v[:GIOI_HAN_O]]) + f"\n… cắt bớt {bo} dòng (Excel giới hạn 32.767 ký tự/ô)"
    return v


def doc_doi_chieu(xlsx):
    """Đọc lại file doi_chieu_*.xlsx (cả 2 sheet) -> list dict cùng khoá COT_KQ, sắp theo (năm, dòng)."""
    import openpyxl
    wb = openpyxl.load_workbook(xlsx, read_only=True)
    ten2k = {h: k for k, h in COT_KQ}
    kq = []
    for ws in (wb[n] for n in ("KHOP", "CAN_XU_LY") if n in wb.sheetnames):
        rows = ws.iter_rows(values_only=True)
        keys = [ten2k.get(h) for h in next(rows, ())]
        for row in rows:
            r = {k: ("" if v is None else v) for k, v in zip(keys, row) if k}
            if r.get("ten"):
                kq.append(r)
    wb.close()
    kq.sort(key=lambda r: (r["nam"], r["dong"]))
    return kq


# ---------- kiểm tra nội dung ----------
def loi_noi_dau(text):
    """Trong 5000 ký tự đầu của 'Lời nói đầu' / 'Lời Nhà xuất bản' / 'Lời giới thiệu', chỉ giữ các đoạn bắt đầu bằng
    'cuốn sách / bộ sách / ...' (dài 600 ký tự) — đủ để dò tên sách, cache nhỏ. '' nếu không có."""
    t = chuan_hoa(text)
    m = re.search(r"\bloi (noi dau|nha xuat ban|gioi thieu|tua|mo dau)\b", t)
    t = t[m.start():m.start() + 5000] if m else ""
    return " | ".join(t[x.start():x.start() + 600] for x in re.finditer(NHAC_SACH, t))


def nam_trang_cuoi(t):
    """Năm 20xx ở trang cuối (text đã chuan_hoa). Có khối thông tin xuất bản ('đăng ký / quyết định / giấy phép
    xuất bản', 'nộp lưu chiểu', cả font lỗi 'noap lou chieau') thì chỉ lấy năm từ đó trở đi, để bỏ năm của luật,
    tài liệu tham khảo...; không có khối này thì không lấy năm nào (trang cuối có thể là mục lục '... (1900 - 2020)')."""
    m = re.search(r"\b(quyet dinh|giay phep|dang ky)( ke hoach)? xuat ban\b|\bl\w{0,2}u chi\w{0,3}u\b", t)
    return sorted({int(y) for y in re.findall(r"\b20\d\d\b", t[m.start():])}) if m else []


def nam_trang_dau(t):
    return nam_trang_dau_bac(t)[0]


def nam_trang_dau_bac(t):
    """-> (năm, bậc tin cậy 1/2/3, 0 nếu không có). Năm xuất bản ở 8 trang đầu (text đã chuan_hoa), lấy theo thứ tự tin cậy, có bậc nào thì dừng ở bậc đó:
    1. khối thông tin XB đặt ở đầu sách (đăng ký / quyết định XB, lưu chiểu), tính đến chữ ISBN;
    2. biên mục CIP '... quốc gia, 2021. - 312tr.' và trang tên sách 'Nhà xuất bản ... Hà Nội - 2021';
    3. chữ ký lời NXB 'tháng 3 năm 2021 Nhà xuất bản' (có thể ký trước năm XB).
    Bỏ qua năm trong nội dung và chú thích tài liệu tham khảo ('Nxb. ..., Hà Nội, 2011, t.7')."""
    nam = lambda x: sorted({int(y) for y in re.findall(r"\b20\d\d\b", x)})
    m = re.search(r"\b(quyet dinh|giay phep|dang ky)( ke hoach)? xuat ban\b|\bl\w{0,2}u chi\w{0,3}u\b", t)
    if m and nam(t[m.start():m.start() + 300].split(" isbn ")[0]):
        return nam(t[m.start():m.start() + 300].split(" isbn ")[0]), 1
    ys = {int(y) for pat in (r"\b(20\d\d) \d+ ?tr\b", r"\bnha xuat ban [a-z ]{0,50}?\bha noi (20\d\d)\b")
          for y in re.findall(pat, t)}
    if ys:
        return sorted(ys), 2
    ys = sorted({int(y) for y in re.findall(r"\bthang \d{1,2} nam (20\d\d) nh\w{0,2} xu", t)})
    return ys, 3 if ys else 0


def khop_nam(t, nam):
    """Ưu tiên trang cuối, không khớp thì xét 8 trang đầu -> 'trang cuối' / '8 trang đầu' / ''.
    Ngày ký lời NXB (bậc 3) thường là năm in lần đầu, sách tái bản vẫn giữ: không dùng khi trang cuối đã ghi năm khác."""
    if nam in t["nam"]:
        return "trang cuối"
    if nam in t.get("dau", ()) and not (t["nam"] and t.get("dau_bac") == 3):
        return "8 trang đầu"
    return ""


def cac_nam(t):
    return sorted(set(t["nam"]) | set(t.get("dau", ())))


_WORD = [None, None]  # [Word.Application dùng chung (False = máy không có Word), PID của tiến trình Word đó]
WORD_HET_GIO = 90  # giây; quá thời gian (file hỏng, hộp thoại treo...) thì tắt Word, bỏ qua file đó
_KHONG_CUA_SO = 0x08000000  # CREATE_NO_WINDOW: không bật cửa sổ đen khi gọi tasklist/taskkill


def _pid_word():
    out = subprocess.run(["tasklist", "/FI", "IMAGENAME eq WINWORD.EXE", "/FO", "CSV", "/NH"],
                         capture_output=True, text=True, creationflags=_KHONG_CUA_SO).stdout
    return {int(x) for x in re.findall(r'"WINWORD\.EXE","(\d+)"', out, re.I)}


def _word_app():
    """Word.Application dùng chung (mở lần đầu cần); None nếu máy không có Word."""
    if _WORD[0] is None:
        try:
            import pythoncom, win32com.client
            pythoncom.CoInitialize()
            truoc = _pid_word()
            _WORD[0] = win32com.client.DispatchEx("Word.Application")
            _WORD[0].Visible, _WORD[0].DisplayAlerts = False, 0
            _WORD[1] = max(_pid_word() - truoc, default=None)  # Word riêng của tool, không đụng Word người dùng đang mở
        except Exception:
            _WORD[0] = False
    return _WORD[0] or None


def _giet_word():
    if _WORD[1]:
        subprocess.run(["taskkill", "/PID", str(_WORD[1]), "/F"], capture_output=True, creationflags=_KHONG_CUA_SO)


def _doc_word(path):
    """.doc -> (KY_DAU ký tự đầu, KY_CUOI ký tự cuối) qua Microsoft Word, mở chỉ đọc; None nếu không đọc được.
    Lấy theo số ký tự chứ không đếm trang: đếm trang bắt Word dàn trang cả cuốn, chậm ~20 lần."""
    with open(duong_dan_dai(path), "rb") as f:
        if f.read(8) != b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":  # không phải file Word 97-2003 thật
            return None
    if os.name != "nt":
        return _doc_linux(path)
    if not _word_app():
        return None
    mo, tam = os.path.abspath(path), None
    if len(mo) > 200:  # Word không mở được đường dẫn dài hơn ~255 ký tự (trả về rỗng, không báo lỗi) -> mở bản copy ngắn
        fd, tam = tempfile.mkstemp(suffix=".doc", prefix="sxb_"); os.close(fd)
        shutil.copyfile(duong_dan_dai(path), tam); mo = tam
    bi_giet = []
    hen = threading.Timer(WORD_HET_GIO, lambda: (bi_giet.append(1), _giet_word())); hen.start()
    d = None
    try:
        # mật khẩu giả: file có mật khẩu báo lỗi ngay thay vì hiện hộp thoại chờ người nhập
        d = _WORD[0].Documents.Open(mo, ConfirmConversions=False, ReadOnly=True, AddToRecentFiles=False,
                                    PasswordDocument="sxb", Visible=False, NoEncodingDialog=True)
        if d is None:
            raise OSError("Word không mở được file")
        het = d.Content.End
        return d.Range(0, min(het, KY_DAU)).Text, d.Range(max(0, het - KY_CUOI), het).Text
    except Exception:
        try:
            _WORD[0].Version  # Word còn sống? (có thể bị tắt vì quá giờ hoặc tự chết)
        except Exception:
            bi_giet.append(1)
        raise
    finally:
        hen.cancel()
        try:
            d is not None and d.Close(False)
        except Exception:
            pass
        if bi_giet:  # Word đã bị tắt (quá giờ) hoặc tự chết -> lần sau mở Word mới
            _WORD[0] = _WORD[1] = None
        if tam:
            try:
                os.remove(tam)
            except OSError:
                pass


def _doc_linux(path):
    """.doc trên Linux (không có Word): antiword, không có thì catdoc (sudo apt install antiword catdoc).
    -> (KY_DAU ký tự đầu, KY_CUOI ký tự cuối); None nếu máy không có hai lệnh này hoặc đọc lỗi/quá giờ."""
    for lenh in (["antiword", "-m", "UTF-8.txt", "-w", "0"], ["catdoc", "-d", "utf-8", "-w"]):
        if shutil.which(lenh[0]):
            try:
                r = subprocess.run(lenh + [path], capture_output=True, timeout=WORD_HET_GIO)
            except subprocess.TimeoutExpired:
                return None
            t = r.stdout.decode("utf-8", "replace")
            if r.returncode == 0 and t.strip():
                return t[:KY_DAU], t[-KY_CUOI:]
    return None


def _docx_text(path):
    """Chữ của .docx đọc thẳng từ word/document.xml (nhanh hơn python-docx nhiều lần), mỗi đoạn một dòng."""
    with zipfile.ZipFile(path) as z:
        x = z.read("word/document.xml").decode("utf-8")
    x = re.sub(r"<w:(?:tab|br|cr)\b[^>]*/>", "<w:t> </w:t>", x)  # tab/xuống dòng giữa các chữ -> khoảng trắng
    doan = (html.unescape("".join(re.findall(r"<w:t(?:\s[^>]*)?>([^<]*)</w:t>", d))) for d in x.split("</w:p>"))
    return "\n".join(d for d in doan if d.strip())


def dong_word():
    """Tắt Word sau khi chạy tab 3."""
    if _WORD[0]:
        try:
            _WORD[0].Quit(False)
        except Exception:
            _giet_word()
    _WORD[0] = _WORD[1] = None


def trich(path, so_trang=8):
    """-> {'nam': năm XB ở trang cuối, 'dau': năm XB ở 8 trang đầu, 'lnd': đoạn 'cuốn sách ...' trong lời nói đầu}.
    PDF theo trang; .docx/.doc theo số ký tự KY_DAU/KY_CUOI (.doc đọc qua Microsoft Word). None nếu không đọc được."""
    ext = os.path.splitext(path)[1].lower()
    p, them = duong_dan_dai(path), {}
    try:
        if ext == ".pdf":
            import fitz
            with fitz.open(p) as d:
                n = len(d)
                xem = sorted(set(range(min(so_trang, n))) | set(range(max(0, n - 2), n)))
                chu = {i: d[i].get_text() for i in xem}
                dau = " ".join(chu[i] for i in range(min(so_trang, n)))
                cuoi = " ".join(chu[i] for i in range(max(0, n - 2), n))
                # ponytail: không có OCR - PDF gần như không có lớp chữ thì chỉ đánh dấu 'cần OCR' (sheet CAN_OCR);
                # muốn đọc được thì OCR riêng các trang này (vd Tesseract + gói tiếng Việt) rồi lưu cache như file khác
                ocr = ""
                if n and len(re.sub(r"\s", "", "".join(chu.values()))) < 100:
                    ocr = "scan" if any(d[i].get_images() for i in xem) else "hinh_ve"
                them = dict(ocr=ocr, trang=n)
        elif ext == ".docx":
            t = _docx_text(p)
            dau, cuoi = t[:KY_DAU], t[-KY_CUOI:]
        elif ext == ".doc":
            r = _doc_word(path)
            if r is None:
                return None
            dau, cuoi = r
        else:
            return None
    except Exception:
        return None
    nam_dau, bac = nam_trang_dau_bac(chuan_hoa(dau))
    return dict(nam=nam_trang_cuoi(chuan_hoa(cuoi)), dau=nam_dau, dau_bac=bac, lnd=loi_noi_dau(dau), **them)


def _pid_song(pid):
    """Tiến trình pid còn chạy không."""
    if os.name != "nt":
        try:
            os.kill(pid, 0)
        except ProcessLookupError:
            return False
        except PermissionError:
            pass
        return True
    import ctypes
    k = ctypes.windll.kernel32
    h = k.OpenProcess(0x00100000, False, pid)  # SYNCHRONIZE
    if not h:
        return False
    try:
        return k.WaitForSingleObject(h, 0) == 0x102  # WAIT_TIMEOUT: chưa kết thúc
    finally:
        k.CloseHandle(h)


class Cache:
    """Cache kết quả đọc file trong SQLite trên đĩa: không chiếm RAM, lưu ngay sau từng file, chạy lại dùng tiếp.
    Trước khi đọc mỗi file ghi dấu 'đang đọc'. Lần mở sau còn dấu của một tiến trình đã chết = tool tắt đột ngột khi
    đọc file đó (thường do file PDF hỏng làm thư viện đọc PDF sập) -> file đó coi là không đọc được, bỏ qua, ghi vào vua_loi."""
    def __init__(self, path):
        self.db = sqlite3.connect(path, check_same_thread=False)
        self.db.execute("pragma journal_mode=wal")  # ghi nhanh, lưu sau từng file vẫn không chậm
        self.db.execute("pragma synchronous=normal")
        self.db.execute("create table if not exists c (k text primary key, v text)")
        self.db.execute("create table if not exists dang_doc (pid integer primary key, k text, path text)")
        self.db.execute("create table if not exists file_loi (path text primary key, ly_do text, luc text)")
        self.vua_loi = []
        for pid, k, p in self.db.execute("select pid, k, path from dang_doc").fetchall():
            if pid == os.getpid() or _pid_song(pid):
                continue  # đang có tiến trình khác / luồng khác đọc thật
            self.loi(k, p, "tool tắt đột ngột khi đang đọc file này")
            self.db.execute("delete from dang_doc where pid = ?", (pid,))
        self.db.commit()

    def loi(self, k, path, ly_do):
        """File làm sập tool / tiến trình đọc: lưu là không đọc được (lần sau bỏ qua), ghi vào file_loi và vua_loi."""
        self[k] = None
        self.db.execute("insert or replace into file_loi values (?, ?, ?)", (path, ly_do, dt.datetime.now().isoformat(" ", "seconds")))
        self.vua_loi.append(path)

    def doc(self, k, path, ham):
        """Đọc 1 file: ghi dấu 'đang đọc' xuống đĩa, đọc, lưu kết quả, xoá dấu."""
        self.db.execute("insert or replace into dang_doc values (?, ?, ?)", (os.getpid(), k, path)); self.db.commit()
        v = ham()
        self[k] = v
        self.db.execute("delete from dang_doc where pid = ?", (os.getpid(),)); self.db.commit()
        return v

    def get(self, k, mac_dinh=None):
        row = self.db.execute("select v from c where k = ?", (k,)).fetchone()
        return json.loads(row[0]) if row else mac_dinh

    def __contains__(self, k):
        return self.get(k, KeyError) is not KeyError

    def __getitem__(self, k):
        v = self.get(k, KeyError)
        if v is KeyError:
            raise KeyError(k)
        return v

    def __setitem__(self, k, v):
        self.db.execute("insert or replace into c values (?, ?)", (k, json.dumps(v, ensure_ascii=False)))

    def __len__(self):
        return self.db.execute("select count(*) from c").fetchone()[0]

    def dong(self):
        self.db.commit(); self.db.close()


def doc_cache(path):
    return Cache(path)


def ghi_cache(cache, path):
    """Lưu cache xuống đĩa (Cache: commit; dict thường: chép sang file SQLite)."""
    if isinstance(cache, Cache):
        cache.db.commit()
    else:
        c = Cache(path)
        for k, v in cache.items():
            c[k] = v
        c.dong()


def _khoa_cache(path):
    """Khoá cache theo (path, mtime); None nếu file không còn. |10: đổi khi cách trích thay đổi."""
    try:
        return f"{path}|{int(os.path.getmtime(duong_dan_dai(path)))}|10"
    except OSError:
        return None


def trich_cache(cache, path):
    """Như trich() nhưng nhớ theo (path, mtime) trong cache dict."""
    k = _khoa_cache(path)
    if k is None:
        return None
    v = cache.get(k, KeyError)
    if v is KeyError:
        if isinstance(cache, Cache):
            v = cache.doc(k, path, lambda: trich(path))
        else:
            v = cache[k] = trich(path)
    return v


def doc_truoc(cache, paths, so_luong=SO_LUONG, progress=None, ham=trich):
    """Đọc trước song song (so_luong tiến trình con) các file PDF/DOCX chưa có trong cache (Cache), lưu vào cache để
    các bước sau lấy ra ngay. .doc phải đọc qua Word nên vẫn đọc lần lượt ở bước sau.
    Tiến trình con sập (PDF hỏng làm thư viện đọc PDF sập) không làm tắt tool, chỉ làm hỏng các file đang chờ trong lượt đó:
    đọc lại từng file này trong tiến trình riêng, file nào làm sập lần nữa thì cache.loi() và coi là không đọc được.
    progress(xong, tong). Bấm Dừng: DaDung, phần đã đọc giữ trong cache. -> số file đã đọc."""
    from concurrent.futures import FIRST_COMPLETED, ProcessPoolExecutor, wait
    from concurrent.futures.process import BrokenProcessPool
    import multiprocessing
    mp = multiprocessing.get_context("spawn")  # Linux mặc định fork: fork từ tiến trình đang có luồng giao diện dễ treo
    viec = {}
    for p in paths:
        k = p.lower().endswith((".pdf", ".docx")) and _khoa_cache(p)
        if k and k not in viec and k not in cache:
            viec[k] = p
    con, nghi, tong, xong = list(viec.items())[::-1], [], len(viec), 0

    def ghi(kp, v):
        nonlocal xong
        cache[kp[0]] = v; xong += 1
        if xong % 200 == 0:
            cache.db.commit()
        if progress:
            progress(xong, tong)

    # ponytail: không có thời hạn cho từng file - một PDF làm thư viện treo hẳn (hiếm) sẽ giữ một tiến trình tới khi bấm Dừng
    while con:
        ex, dang, hong = ProcessPoolExecutor(so_luong, mp_context=mp), {}, False
        try:
            while dang or (con and not hong):
                while con and not hong and len(dang) < so_luong * 4:  # giữ hàng chờ ngắn: ít file phải đọc lại khi sập
                    kp = con.pop()
                    try:
                        dang[ex.submit(ham, kp[1])] = kp
                    except BrokenProcessPool:
                        con.append(kp); hong = True
                _dung()
                for f in wait(dang, timeout=0.5, return_when=FIRST_COMPLETED)[0]:
                    kp = dang.pop(f)
                    try:
                        ghi(kp, f.result())
                    except BrokenProcessPool:
                        nghi.append(kp); hong = True
        finally:
            ex.shutdown(wait=True, cancel_futures=True)
    for kp in nghi:  # lượt có tiến trình sập: đọc lại riêng từng file để tìm đúng file gây sập
        _dung()
        with ProcessPoolExecutor(1, mp_context=mp) as ex:
            try:
                v = ex.submit(ham, kp[1]).result()
            except BrokenProcessPool:
                cache.loi(kp[0], kp[1], "file làm sập tiến trình đọc PDF (file hỏng?)"); xong += 1
                continue
        ghi(kp, v)
    cache.db.commit()
    return tong


def _trich_thu(path):
    """Chỉ dùng trong _test: file tên 'sap' làm tiến trình con chết đột ngột như PDF hỏng làm sập thư viện."""
    if "sap" in os.path.basename(path):
        os._exit(3)
    return trich(path)


def _dong_file(cac_file):
    """'Tất cả file tìm được' -> [(date_modified, path)]"""
    return [(l[:19], l[21:]) for l in str(cac_file or "").splitlines() if len(l) > 21 and l[:4].isdigit()]


def nam_ten_file(path, key=""):
    """Năm 20xx trong tên file ('... (TB12-4-2021).pdf', '... 2021.pdf'), trừ năm nằm trong chính tên sách."""
    ten = os.path.splitext(os.path.basename(path))[0]
    return {int(y) for y in re.findall(r"(?<!\d)20\d\d(?!\d)", ten)} - {int(y) for y in re.findall(r"\b20\d\d\b", key)}


def _khong_lech(cache, p, nam, key=""):
    """File cùng thư mục được nhận khi không ghi năm khác: cả trong nội dung lẫn trong tên file."""
    fy, t = nam_ten_file(p, key), trich_cache(cache, p)
    return (not fy or nam in fy) and (not t or not cac_nam(t) or bool(khop_nam(t, nam)))


def _chon(r, ok, ghi_chu, cung_tm=()):
    """ok: [(date_modified, path)] đã xác nhận năm -> ghi KHỚP vào r. File trong cung_tm nằm cùng thư mục
    với file đã xác nhận cũng được nhận (vd .doc cạnh PDF). False nếu ok rỗng."""
    if not ok:
        return False
    tm, key = {os.path.dirname(p) for _, p in ok}, chuan_hoa(r["ten"])
    w, p, _ = chon_file(ok, key)
    w2, p2, _ = chon_file([x for x in cung_tm if os.path.dirname(x[1]) in tm], key)
    w, p = w or w2, p or p2
    r.update(tinh_trang="KHỚP", file_word=w, file_pdf=p, ghi_chu=ghi_chu if w or p else ghi_chu + "; chỉ có file bìa")
    return True


def file_theo_dm(kq):
    """{đường dẫn: năm} các file đã khớp sách theo Date Modified ở bước khớp tên (dòng KHỚP, file cùng năm với dòng)."""
    return {p: r["nam"] for r in kq if r["tinh_trang"] == "KHỚP"
            for d, p in _dong_file(r["cac_file"]) if d[:4] == str(r["nam"])}


def xu_ly_lech_nam(r, cache, theo_dm=None):
    """Dòng LỆCH NĂM: mở từng file trong 'Tất cả file tìm được', nhận file có năm cột Năm ở trang cuối,
    không có thì ở 8 trang đầu (cùng các file khác trong thư mục đó).
    theo_dm (file_theo_dm): file đã thuộc sách năm khác theo Date Modified thì không xét lại bằng nội dung."""
    theo_dm = theo_dm or {}
    ds = [x for x in _dong_file(r["cac_file"]) if theo_dm.get(x[1], r["nam"]) == r["nam"]]
    ok, ghi, thay, nguon, ocr = [], [], set(), set(), []
    for _, p in _dong_file(r["cac_file"]):
        if theo_dm.get(p, r["nam"]) != r["nam"]:
            ghi.append(f"{os.path.basename(p)} đã thuộc sách năm {theo_dm[p]} theo Date Modified")
    for d, p in ds:
        if not p.lower().endswith((".pdf", ".doc", ".docx")):  # ảnh bìa... không có năm để đọc
            continue
        t = trich_cache(cache, p)
        if t is None:
            ghi.append(f"bỏ qua (không đọc được): {os.path.basename(p)}")
            continue
        thay.update(cac_nam(t))
        if t.get("ocr") and not la_bia(p, chuan_hoa(r["ten"])) and not pdf_mong(p):
            ocr.append(f"{os.path.basename(p)} ({t['trang']} trang)")
        if khop_nam(t, r["nam"]):
            ok.append((d, p)); nguon.add(khop_nam(t, r["nam"]))
    nhan = f"khớp theo nội dung ({', '.join(sorted(nguon, reverse=True))})"
    if not _chon(r, ok, nhan, [x for x in ds if _khong_lech(cache, x[1], r["nam"], chuan_hoa(r["ten"]))]):
        ghi.insert(0, f"không thấy năm {r['nam']} ở trang cuối/8 trang đầu (thấy: {', '.join(map(str, sorted(thay))) or '-'})")
        if ocr:
            ghi.append("cần OCR (PDF không có lớp chữ): " + ", ".join(ocr))
        r["ghi_chu"] = "; ".join(ghi)


def loc_files(files, thu_muc="", nam=""):
    """Giới hạn file kho: trong thư mục con, Date Modified trong năm/khoảng năm ('2015' hoặc '2015-2017')."""
    tm = os.path.normcase(os.path.abspath(thu_muc)) if thu_muc.strip() else ""
    ns = [int(x) for x in re.findall(r"\d{4}", nam)]
    lo, hi = (ns[0], ns[-1]) if ns else (0, 9999)
    return [f for f in files if lo <= f["nam"] <= hi
            and (not tm or os.path.normcase(f["path"]).startswith(tm))]


NHAC_SACH = r"\b(?:cuon sach|cuon|bo sach|tap sach|an pham)(?: (?:chuyen khao|tham khao|nay))? "


def nhac_ten_sach(lnd, key):
    """Lời nói đầu giới thiệu đúng cuốn này: '... xuất bản cuốn sách <tên sách>' — không chỉ nhắc qua tên sách khác."""
    return re.search(NHAC_SACH + re.escape(key) + r"\b", lnd) is not None


def xu_ly_khong_tim_thay(rows, files, cache, progress=None):
    """Dòng KHÔNG TÌM THẤY: tìm trong lời nói đầu câu giới thiệu 'cuốn sách <tên sách>', rồi kiểm tra năm trang cuối."""
    theo_prefix = defaultdict(list)
    for r in rows:
        r["_key"], r["_tim"] = loi_ten(r["ten"]), {}
        theo_prefix[" ".join(r["_key"].split()[:3])].append(r)
    theo_tm = defaultdict(list)
    for f in files:
        theo_tm[os.path.dirname(f["path"])].append((f["date_modified"], f["path"]))
    for i, f in enumerate(files, 1):
        _dung()
        t = f["dinh_dang"] in DOC_DUOC and trich_cache(cache, f["path"])
        if t and t["lnd"]:
            w, lnd = t["lnd"].split(), f" {t['lnd']} "
            for j in range(len(w)):  # ponytail: tra từng cụm 1-3 từ của lời nói đầu vào chỉ mục tên sách
                for n in (1, 2, 3):
                    for r in theo_prefix.get(" ".join(w[j:j + n]), ()):
                        if nhac_ten_sach(lnd, r["_key"]):
                            r["_tim"][f["path"]] = (f, t)
        if progress:
            progress(i)
    for r in rows:
        tim, _ = r.pop("_tim"), r.pop("_key")
        if not tim:
            r["tinh_trang"] = "KHÔNG TÌM ĐƯỢC"
            continue
        r.update(so_file=len(tim), kieu_khop="lời nói đầu",
                 cac_file=_ds_file((f["date_modified"], f["path"]) for f, _ in tim.values()))
        ok = [(f["date_modified"], f["path"]) for f, t in tim.values() if khop_nam(t, r["nam"])]
        cung = [x for _, p in ok for x in theo_tm[os.path.dirname(p)] if _khong_lech(cache, x[1], r["nam"], chuan_hoa(r["ten"]))]
        if not _chon(r, ok, "khớp theo nội dung (lời nói đầu)", cung):
            r["tinh_trang"] = "LỆCH NĂM - tìm theo lời nói đầu"
            r["ghi_chu"] = f"không thấy năm {r['nam']} ở trang cuối/8 trang đầu"


# ---------- đổi tên & lưu ----------
def ke_hoach_luu(kq, thu_muc_dich):
    """-> [(r, nguồn, đích)] cho dòng KHỚP. Mỗi sách một thư mục:
    <đích>/<Năm>/thang <MM>/<Tên file mới>/<Tên file mới><đuôi>; thư mục trùng tên thêm _2, _3.
    File bìa (ds_bia) vào <thư mục sách>/Bia/<tên mới đổi _R_ thành _B_><đuôi>; nhiều bìa cùng đuôi thì _2, _3,
    tên đã có file khác thì lấy tên kế tiếp (không ghi đè)."""
    da_dung, kh = set(), []
    for r in kq:
        if r["tinh_trang"] != "KHỚP":
            continue
        nguon, cac_bia = [x for x in (r["file_word"], r["file_pdf"]) if x], ds_bia(r)
        if not (nguon or cac_bia):  # sách chỉ có bìa: vẫn lưu bìa vào <thư mục sách>/Bia
            continue
        goc = os.path.join(thu_muc_dich, str(r["nam"]), f"thang {thang_qdxb(r['qdxb'])}", r["ten_moi"])
        tm, i = goc, 1
        # thư mục đã có: dùng lại nếu là của chính sách này (chạy lại sau khi dừng giữa chừng), khác thì sang _2, _3
        while tm.lower() in da_dung or (os.path.exists(duong_dan_dai(tm)) and not all(
                da_co(src, os.path.join(tm, os.path.basename(tm) + os.path.splitext(src)[1].lower()), True) for src in nguon)):
            i += 1
            tm = f"{goc}_{i}"
        da_dung.add(tm.lower())
        ten = os.path.basename(tm)
        kh += [(r, src, os.path.join(tm, ten + os.path.splitext(src)[1].lower())) for src in nguon]
        bia = ten_bia(r["ten_moi"])
        for src in cac_bia:
            duoi, j = os.path.splitext(src)[1].lower(), 1
            dst = os.path.join(tm, "Bia", bia + duoi)
            while dst.lower() in da_dung or (os.path.exists(duong_dan_dai(dst)) and not da_co(src, dst)):
                j += 1
                dst = os.path.join(tm, "Bia", f"{bia}_{j}{duoi}")
            da_dung.add(dst.lower())
            kh.append((r, src, dst))
    return kh


def da_co(src, dst, chua_co_cung_duoc=False):
    """File đích đã là bản của nguồn (cùng dung lượng, hoặc nguồn đã bị di chuyển đi).
    chua_co_cung_duoc: đích chưa có cũng tính là được (thư mục copy dở)."""
    s, d = duong_dan_dai(src), duong_dan_dai(dst)
    if not os.path.exists(d):
        return chua_co_cung_duoc
    return not os.path.exists(s) or os.path.getsize(s) == os.path.getsize(d)


def _con_trong(p):
    while not os.path.exists(duong_dan_dai(p)) and os.path.dirname(p) != p:
        p = os.path.dirname(p)
    return shutil.disk_usage(duong_dan_dai(p)).free


def luu_file(kh, di_chuyen=False, progress=None):
    """Copy (hoặc move) theo kế hoạch, ghi đường dẫn đích vào r['dich']. File đã có thì bỏ qua (chạy lại được);
    lỗi ở file nào thì ghi vào Ghi chú rồi làm tiếp. -> (số file đã chép, số đã có sẵn, [lỗi]).
    progress(i, byte đã chép, tổng byte cần chép) sau mỗi file: thời gian chép theo dung lượng, không theo số file."""
    can = [(r, src, p) for r, src, p in kh if not da_co(src, p)]
    cung_o = lambda src, p: os.path.splitdrive(os.path.abspath(src))[0].lower() == os.path.splitdrive(os.path.abspath(p))[0].lower()
    kt = {(src, p): os.path.getsize(duong_dan_dai(src)) for _, src, p in can if os.path.exists(duong_dan_dai(src))}  # theo cặp: 1 nguồn có thể chép tới 2 nơi
    byte = sum(kt.get((src, p), 0) for _, src, p in can if not (di_chuyen and cung_o(src, p)))
    if can and byte > _con_trong(can[0][2]):
        raise ValueError(f"Ổ đích không đủ chỗ: cần {byte / 2**30:.1f} GB, còn {_con_trong(can[0][2]) / 2**30:.1f} GB")
    chep, co_san, loi, da_chep, tong = 0, 0, [], 0, sum(kt.values())
    for i, (r, src, p) in enumerate(kh, 1):
        if DUNG.is_set():  # dừng: phần đã chép giữ nguyên, chạy lại sẽ làm tiếp
            break
        try:
            if da_co(src, p):
                co_san += 1
            else:
                os.makedirs(duong_dan_dai(os.path.dirname(p)), exist_ok=True)
                tam = p + ".dang_chep"  # chép ra tên tạm rồi đổi tên: dừng giữa chừng không để lại file dở mang tên thật
                (shutil.move if di_chuyen else shutil.copy2)(duong_dan_dai(src), duong_dan_dai(tam))
                os.replace(duong_dan_dai(tam), duong_dan_dai(p))
                chep += 1; da_chep += kt.get((src, p), 0)
            r["dich"] = "\n".join(x for x in (r.get("dich", ""), p) if x)
        except OSError as e:
            loi.append(f"{src}: {e}")
            r["ghi_chu"] = "; ".join(x for x in (r.get("ghi_chu", ""), f"LỖI lưu {os.path.basename(src)}: {e.strerror or e}") if x)
        if progress:
            progress(i, da_chep, tong)
    return chep, co_san, loi


def moi_nhat(*mau):
    """File mới nhất (theo giờ sửa) khớp các mẫu glob; '' nếu chưa có."""
    fs = [f for m in mau for f in glob.glob(m)]
    return max(fs, key=os.path.getmtime) if fs else ""


# ---------- giao diện ----------
def gui(so_tab=2):
    """Giao diện. so_tab=2 (SXBTool.exe): 1. Đối chiếu (quét kho -> khớp tên -> kiểm tra nội dung, một nút) | 2. Đổi tên & lưu.
    so_tab=4 (SXBTool_4tab.exe): 1. Quét kho | 2. Đối chiếu Excel | 3. Kiểm tra nội dung | 4. Đổi tên & lưu - từng bước riêng.
    Hai bản dùng chung các bước xử lý bên dưới, chung thư mục ketqua/ và chung cache."""
    import tkinter as tk
    from tkinter import ttk, filedialog, messagebox

    # kết quả gom vào ketqua/ cạnh file exe (hoặc cạnh sxb_tool.py)
    noi = os.path.dirname(sys.executable if getattr(sys, "frozen", False) else os.path.abspath(__file__))
    goc = os.path.join(noi, "ketqua")
    # Excel tổng hợp mặc định: cạnh exe (hoặc sxb_tool.py), không có thì ở thư mục cha (exe nằm trong dist\)
    XL_MD = next((f for d in (noi, os.path.dirname(noi)) for f in [os.path.join(d, XL_TONG_HOP)] if os.path.isfile(f)), "")
    TM = {k: os.path.join(goc, k) for k in ("1_quet_kho", "2_doi_chieu", "3_doi_ten", "sach_da_luu")}
    for d in TM.values():
        os.makedirs(d, exist_ok=True)
    KHO = (os.path.join(TM["1_quet_kho"], "kho_*.xlsx"), os.path.join(TM["1_quet_kho"], "kho_*.csv"))
    # file doi_chieu mới nhất (cả thư mục 3_kiem_tra của phiên bản cũ)
    DC = (os.path.join(TM["2_doi_chieu"], "doi_chieu_*.xlsx"), os.path.join(goc, "3_kiem_tra", "doi_chieu_*.xlsx"))
    # cache đọc nội dung dùng chung mọi lần chạy (khoá theo đường dẫn + ngày sửa file) nên quét lại kho vẫn dùng được
    CACHE = os.path.join(TM["1_quet_kho"], "cache_noi_dung.sqlite")

    # nhật ký lỗi: cả lỗi sập ở tầng C (vd file PDF hỏng làm thư viện đọc PDF sập) để biết tool tắt ở đâu
    NK_LOI = os.path.join(goc, "nhat_ky_loi.txt")
    flog = open(NK_LOI, "a", encoding="utf-8", buffering=1)
    flog.write(f"\n=== mở tool {dt.datetime.now():%d/%m/%Y %H:%M:%S} (bản {so_tab} tab, pid {os.getpid()}) ===\n")
    faulthandler.enable(flog, all_threads=True)
    # quét kho ghi dần vào nhật ký này; tool tắt / bấm Dừng giữa chừng thì lần sau hỏi quét tiếp
    NK_QUET = os.path.join(TM["1_quet_kho"], "quet_dang_do.csv")

    root = tk.Tk()
    root.title("Tool SXB - đối chiếu kho sách với Excel tổng hợp, đổi tên & lưu" + (" (bản 4 tab)" if so_tab == 4 else ""))
    root.geometry("820x660")
    thanh = ttk.Frame(root); thanh.pack(side="bottom", fill="x", padx=8, pady=(0, 8))
    b_dung = ttk.Button(thanh, text="■ Dừng", state="disabled", command=lambda: (DUNG.set(), b_dung.config(state="disabled")))
    b_dung.pack(side="right")
    ttk.Label(thanh, text="Dừng sau file đang xử lý; chạy lại sẽ làm tiếp phần còn lại.").pack(side="right", padx=8)
    nb = ttk.Notebook(root); nb.pack(fill="both", expand=True, padx=8, pady=8)
    ui = lambda f, *a: root.after(0, f, *a)  # mọi thay đổi giao diện từ luồng chạy nền đi qua luồng chính

    def hoi(msg):
        """Hỏi Có/Không từ luồng chạy nền: hộp thoại hiện ở luồng chính, luồng nền chờ câu trả lời."""
        tl, xong = [], threading.Event()
        ui(lambda: (tl.append(messagebox.askyesno("Xác nhận", msg, parent=root)), xong.set()))
        xong.wait()
        return tl[0]
    chon_xlsx = lambda: filedialog.askopenfilename(filetypes=[("Excel", "*.xlsx")])
    chon_csv = lambda: filedialog.askopenfilename(filetypes=[("File kho", "*.xlsx *.csv")])

    def hang(parent, r, nhan, var, chon=None):
        ttk.Label(parent, text=nhan, width=20).grid(row=r, column=0, sticky="w", pady=4)
        ttk.Entry(parent, textvariable=var, width=70).grid(row=r, column=1, sticky="we")
        if chon:
            ttk.Button(parent, text="...", width=3, command=lambda: var.set(chon() or var.get())).grid(row=r, column=2)

    def khung(parent, dong):
        """Thanh tiến độ + dòng trạng thái + khung log, từ hàng `dong` trở xuống."""
        pb = ttk.Progressbar(parent, maximum=1000); pb.grid(row=dong, column=0, columnspan=3, sticky="we", pady=(6, 0))
        lb = ttk.Label(parent, text=""); lb.grid(row=dong + 1, column=0, columnspan=3, sticky="w")
        tx = tk.Text(parent, height=10, wrap="word"); tx.grid(row=dong + 2, column=0, columnspan=3, sticky="we", pady=6)
        return pb, lb, tx

    def bat(pb, ham):
        """Lỗi giữa chừng: dừng thanh chạy qua lại (không đụng thanh đã đầy khi xong)."""
        def chay_ham(*a):
            try:
                ham(*a)
            finally:
                ui(lambda: str(pb.cget("mode")) == "indeterminate" and pb.stop())
        return chay_ham

    dang_chay = [0]

    def chay(nut, ham, txt=None):
        nut["state"] = "disabled"; DUNG.clear(); dang_chay[0] += 1; b_dung.config(state="normal")
        def bao(msg, nhan="LỖI"):
            if isinstance(txt, tk.Text):
                txt.insert("end", f"{nhan}: {msg}\n"); txt.see("end")
            elif txt is not None:
                txt.config(text=f"{nhan}: {msg}")
            if nhan == "LỖI":
                messagebox.showerror("Lỗi", msg, parent=root)
        def xong():
            nut.config(state="normal"); dang_chay[0] -= 1
            if not dang_chay[0]:
                b_dung.config(state="disabled")
        def run():
            try:
                ham()
            except DaDung as e:
                msg = str(e)
                ui(lambda: bao(msg, "ĐÃ DỪNG"))
            except Exception as e:
                msg = str(e) or type(e).__name__
                flog.write(f"{dt.datetime.now():%d/%m/%Y %H:%M:%S} LỖI: {traceback.format_exc()}")
                ui(lambda: bao(msg))
            finally:
                ui(xong)
        threading.Thread(target=run, daemon=True).start()

    def tien_do(pb, lb):
        """-> (bao, het). bao(buoc, xong, tong): thanh tiến độ + '...: 120/690 (17%) - đã chạy 2 phút - còn khoảng 9 phút';
        bao(buoc) không kèm số: thanh chạy qua lại. het(chu): thanh đầy, hiện dòng kết thúc. Vẽ lại tối đa ~3 lần/giây."""
        st = {"buoc": None, "t0": 0.0, "ve": 0.0}
        def bao(buoc, xong=None, tong=None, hien=None):
            bay = time.time()
            if buoc != st["buoc"]:
                st.update(buoc=buoc, t0=bay, ve=0.0)
            elif xong is not None and xong < tong and bay - st["ve"] < 0.3:
                return
            st["ve"] = bay
            if xong is None:
                ui(lambda: (pb.config(mode="indeterminate"), pb.start(15), lb.config(text=f"{buoc}...")))
                return
            da = bay - st["t0"]
            con = f" - còn khoảng {_gio(da / xong * (tong - xong))}" if 0 < xong < tong else ""
            chu = f"{buoc}: {hien or f'{xong}/{tong}'} ({xong / max(tong, 1):.0%}) - đã chạy {_gio(da)}{con}"
            ui(lambda: (pb.stop(), pb.config(mode="determinate", value=1000 * xong / max(tong, 1)), lb.config(text=chu)))
        def het(chu):
            ui(lambda: (pb.stop(), pb.config(mode="determinate", value=1000), lb.config(text=chu)))
        return bao, het

    def logger(txt):
        return lambda s: root.after(0, lambda: (txt.insert("end", s + "\n"), txt.see("end")))

    # ================= các bước dùng chung cho bản 2 tab và 4 tab =================
    def buoc_quet(thu_muc, exts, log, bao, nhan):
        """Quét kho -> file kho_*.xlsx mới trong 1_quet_kho; log số file theo định dạng, ảnh JPG, thư mục không đọc được."""
        kho, bc, t0 = os.path.join(TM["1_quet_kho"], f"kho_{dt.datetime.now():%Y%m%d_%H%M%S}.xlsx"), {}, time.time()
        if os.path.exists(NK_QUET):  # lần quét trước bị dừng hoặc tool tắt giữa chừng
            with open(NK_QUET, encoding="utf-8", newline="") as f:
                dau = next(csv.reader(f), ["#", "?", ""])
            luc = f"{dt.datetime.fromtimestamp(os.path.getmtime(NK_QUET)):%d/%m/%Y %H:%M}"
            if os.path.normcase(dau[1]) == os.path.normcase(os.path.abspath(thu_muc)) and \
                    dau[2] == " ".join(sorted(e if e.startswith(".") else "." + e for e in exts)):
                if not hoi(f"Lần quét kho trước bị dừng hoặc tool tắt giữa chừng (lần ghi cuối {luc}).\n\n"
                           "Có: quét tiếp phần còn lại (bỏ qua các thư mục đã quét).\nKhông: bỏ phần đã quét, quét lại từ đầu."):
                    os.remove(NK_QUET)
            elif hoi(f"Có lần quét dở (lần ghi cuối {luc}) của thư mục khác:\n{dau[1]}\n\nBỏ lần quét dở đó để quét thư mục này?"):
                os.remove(NK_QUET)
            else:
                raise ValueError(f"Chưa quét: đang giữ lần quét dở của {dau[1]} - chọn lại đúng thư mục đó để quét tiếp.")
        bao(nhan)
        def prog(n, root_dir): bao(f"{nhan}: đã quét {n} file, {_gio(time.time() - t0)}")
        try:
            n = quet(thu_muc, exts, kho, prog, bc, NK_QUET)
        except DaDung:
            raise DaDung(f"Đã dừng khi đang quét kho (sau {sum(bc['dinh_dang'].values())} file). Phần đã quét được giữ lại, "
                         "chạy lại sẽ hỏi quét tiếp.")
        if bc.get("tiep"):
            log(f"Quét tiếp lần dở: {bc['tiep']} file đã quét từ trước, không quét lại.")
        log(f"Quét kho: {n} file trong {_gio(time.time() - t0)} → {kho}")
        log("  " + ", ".join(f"{k} {v}" for k, v in bc["dinh_dang"].most_common()))
        anh = bc["dinh_dang"]["JPG"] + bc["dinh_dang"]["JPEG"]
        if n and anh / n > 0.5:
            log(f"  CHÚ Ý: ảnh JPG chiếm {anh / n:.0%} ({anh} file). Nếu đây là ảnh scan từng trang, nên bỏ .jpg .jpeg ở ô "
                "Định dạng rồi chạy lại: file kho nhỏ hơn, các bước sau nhanh và ít RAM hơn.")
        if bc["loi"]:
            ds_loi = kho + ".loi.txt"
            with open(ds_loi, "w", encoding="utf-8-sig") as f:
                f.write("\n".join(bc["loi"]))
            log(f"  CHÚ Ý: {len(bc['loi'])} thư mục/file KHÔNG đọc được (không có quyền, đường dẫn hỏng...) - "
                f"file trong đó không có trong kho. Danh sách: {ds_loi}")
            for x in bc["loi"][:10]:
                log(f"    {x}")
        return kho

    def buoc_noi_dung(kq, files, nhan_lech, nhan_ktt, log, bao):
        """Kiểm tra nội dung: năm xuất bản cho dòng LỆCH NĂM (nhan_lech), lời nói đầu cho dòng KHÔNG TÌM THẤY (nhan_ktt).
        Nhãn rỗng = bỏ bước đó. Sửa thẳng các dòng trong kq."""
        cache = doc_cache(CACHE)
        for x in cache.vua_loi:
            log(f"CHÚ Ý: lần chạy trước tool tắt đột ngột khi đang đọc file này (file hỏng?) - lần này bỏ qua, "
                f"coi là không đọc được: {x}")
            flog.write(f"File làm tool tắt, đã đánh dấu bỏ qua: {x}\n")

        def doc_song_song(paths, nhan):
            """Đọc trước song song các file PDF/DOCX chưa có trong cache; các bước sau lấy từ cache."""
            n0, t0 = len(cache.vua_loi), time.time()
            try:
                n = doc_truoc(cache, paths, progress=lambda x, t: bao(f"{nhan} - đọc song song {SO_LUONG} tiến trình", x, t))
            except DaDung:
                return  # DUNG vẫn bật: bước sau dừng ngay, ghi kết quả phần đã làm
            if n:
                log(f"  Đọc song song {n} file PDF/DOCX ({SO_LUONG} tiến trình): {_gio(time.time() - t0)}")
            for x in cache.vua_loi[n0:]:
                log(f"CHÚ Ý: file làm sập tiến trình đọc PDF (file hỏng?) - bỏ qua, coi là không đọc được: {x}")
                flog.write(f"File làm sập tiến trình đọc, đã đánh dấu bỏ qua: {x}\n")

        try:
            if nhan_lech:
                lech = [r for r in kq if str(r["tinh_trang"]).startswith("LỆCH NĂM")]
                theo_dm = file_theo_dm(kq)  # file đã khớp năm khác theo Date Modified: không đọc nội dung cho dòng này
                # tiến độ tính theo số file (không trùng) phải mở, không theo số dòng: mỗi sách nhiều ít file khác nhau,
                # nhiều dòng (cùng tên sách, khác năm) dùng chung một thư mục -> file đã đọc thì lần sau lấy từ cache
                so = [{p for _, p in _dong_file(r["cac_file"]) if p.lower().endswith((".pdf", ".doc", ".docx"))} for r in lech]
                da, tong, t0 = set(), len(set().union(*so)), time.time()
                log(f"LỆCH NĂM: {len(lech)} dòng, {tong} file Word/PDF cần mở (file đã có trong cache không phải đọc lại)")
                doc_song_song(set().union(*so), nhan_lech)
                bao(nhan_lech, 0, tong)
                for i, (r, n) in enumerate(zip(lech, so), 1):
                    if DUNG.is_set():
                        break
                    xu_ly_lech_nam(r, cache, theo_dm); da |= n; xong = len(da)
                    bao(nhan_lech, xong, tong)
                    if i % 50 == 0:
                        log(f"  {i}/{len(lech)} dòng, {xong}/{tong} file - còn khoảng {_gio((time.time() - t0) / max(xong, 1) * (tong - xong))}")
                        ghi_cache(cache, CACHE)
            if nhan_ktt and not DUNG.is_set():
                ktt = [r for r in kq if str(r["tinh_trang"]).startswith("KHÔNG TÌM")]
                da_khop = {os.path.dirname(x) for r in kq if r["tinh_trang"] == "KHỚP" for x in (r["file_word"], r["file_pdf"]) if x}
                ds = [f for f in files if os.path.dirname(f["path"]) not in da_khop]
                n_doc = sum(f["dinh_dang"] in DOC_DUOC for f in ds)
                uoc = n_doc * GIAY_MOI_FILE / SO_LUONG
                log(f"KHÔNG TÌM THẤY: {len(ktt)} dòng, quét lời nói đầu {n_doc} file PDF/DOCX, ước tính lần đầu ~{_gio(uoc)} "
                    f"với {SO_LUONG} tiến trình (file đã có trong cache nhanh hơn nhiều; ổ chậm thì lâu hơn)")
                if uoc > 1800 and not hoi(
                        f"Quét lời nói đầu {n_doc} file PDF/DOCX ước tính mất khoảng {_gio(uoc)} (ổ cục bộ; ổ mạng có thể lâu hơn).\n\n"
                        "Có thể bấm ■ Dừng bất cứ lúc nào: phần đã đọc được lưu cache, chạy lại sẽ làm tiếp.\n\nVẫn chạy?"):
                    log("  Bỏ qua bước quét lời nói đầu.")
                else:
                    doc_song_song([f["path"] for f in ds if f["dinh_dang"] in DOC_DUOC], nhan_ktt)
                    t0 = time.time()
                    bao(nhan_ktt, 0, len(ds))
                    def prog(i):
                        bao(nhan_ktt, i, len(ds))
                        if i % 200 == 0:
                            log(f"  {i}/{len(ds)} file - còn khoảng {_gio((time.time() - t0) / i * (len(ds) - i))}")
                        if i % 1000 == 0: ghi_cache(cache, CACHE)
                    try:
                        xu_ly_khong_tim_thay(ktt, ds, cache, prog)
                    except DaDung:
                        pass
        finally:
            ghi_cache(cache, CACHE); dong_word(); cache.dong()
        if DUNG.is_set():
            log("ĐÃ DỪNG theo yêu cầu - ghi kết quả phần đã làm; chạy lại sẽ dùng cache, làm tiếp nhanh.")
        DUNG.clear()  # ghi kết quả cho xong dù đã bấm Dừng
        log("Kết quả: " + ", ".join(f"{k}: {v}" for k, v in Counter(r["tinh_trang"] for r in kq).items()))

    def buoc_ghi(kq, files, xl, ra, log, bao):
        """Sheet thư mục kho + danh sách cần OCR (từ cache) -> doi_chieu_*.xlsx và Excel tổng hợp đánh dấu (KHỚP tô vàng)."""
        bao("Lập sheet thư mục kho và danh sách cần OCR")
        cache = doc_cache(CACHE)
        tm, ocr = doi_chieu_thu_muc(kq, files, cache), ds_can_ocr(kq, files, cache); cache.dong()
        for k, v in Counter(x["tinh_trang"] for x in tm).items(): log(f"  Thư mục kho - {k}: {v}")
        if ocr:
            log(f"  Cần OCR: {len(ocr)} file PDF không có lớp chữ (scan / chữ dạng hình vẽ) - danh sách ở sheet CAN_OCR")
        bao("Ghi file kết quả")
        ra1, ra2 = ghi_ket_qua(kq, xl, ra, tm, ocr)
        log(f"File kết quả: {ra1}\nExcel tổng hợp đánh dấu (dòng KHỚP tô vàng): {ra2}")
        ui(v_dc4.set, ra1)  # tab đổi tên & lưu lấy luôn kết quả mới
        return ra1

    def khop_ten(xl, kho, log, bao, nhan):
        """Đọc Excel tổng hợp + file kho, khớp tên sách -> (kq, files)."""
        bao("Đọc file Excel tổng hợp")
        sach = doc_excel(xl); log(f"Excel tổng hợp: {len(sach)} đầu sách")
        # mã BBT quy định: cạnh Excel tổng hợp, cạnh exe hoặc thư mục cha; đọc lại mỗi lần chạy (sửa file không cần mở lại tool)
        f = next((p for d in (os.path.dirname(xl), noi, os.path.dirname(noi)) for p in [os.path.join(d, XL_MA_BBT)] if os.path.isfile(p)), "")
        MA_BBT.clear(); MA_BBT.update(doc_ma_bbt(f) if f else {})
        log(f"Mã BBT: {len(MA_BBT)} mã quy định từ {f}" if MA_BBT else
            f"Mã BBT: chưa có mã quy định ({XL_MA_BBT} {'chưa điền cột F' if f else 'không thấy'}) - dùng BBT bỏ dấu")
        bao(nhan)
        files = doc_kho(kho); kq = doi_chieu(sach, files)
        log(f"Khớp tên: {len(files)} file kho - " + ", ".join(f"{k}: {v}" for k, v in Counter(r["tinh_trang"] for r in kq).items()))
        return kq, files

    NHAN_LECH, NHAN_KTT = "Năm xuất bản (LỆCH NĂM)", "Lời nói đầu (KHÔNG TÌM THẤY)"
    o_excel = []  # các ô "File Excel tổng hợp": tab đổi tên để trống thì lấy ô đã điền ở tab trước

    def hai_hop_kiem(parent, dong):
        v_lech, v_ktt = tk.BooleanVar(value=True), tk.BooleanVar(value=True)
        ttk.Checkbutton(parent, text="Kiểm tra nội dung sách LỆCH NĂM (đọc năm xuất bản ở trang cuối / 8 trang đầu)", variable=v_lech)\
            .grid(row=dong, column=0, columnspan=3, sticky="w")
        ttk.Checkbutton(parent, text="Tìm sách KHÔNG TÌM THẤY theo lời nói đầu (đọc các file còn lại trong kho - lâu)", variable=v_ktt)\
            .grid(row=dong + 1, column=0, columnspan=3, sticky="w")
        return v_lech, v_ktt

    if so_tab == 2:
        # --- tab 1: đối chiếu = quét kho -> khớp tên sách -> kiểm tra nội dung -> ghi kết quả ---
        t1 = ttk.Frame(nb, padding=10); nb.add(t1, text="1. Đối chiếu")
        v_tm, v_ext, v_kho, v_xl = tk.StringVar(), tk.StringVar(value=EXT_MAC_DINH), tk.StringVar(), tk.StringVar(value=XL_MD)
        v_ra = tk.StringVar(value=TM["2_doi_chieu"]); o_excel.append(v_xl)
        hang(t1, 0, "Thư mục kho", v_tm, filedialog.askdirectory)
        hang(t1, 1, "Định dạng (cách)", v_ext)
        hang(t1, 2, "File kho có sẵn", v_kho, chon_csv)
        ttk.Label(t1, text="(để trống = quét kho mới; chọn file kho_*.xlsx đã quét để bỏ qua bước quét)").grid(row=3, column=1, sticky="w")
        hang(t1, 4, "File Excel tổng hợp", v_xl, chon_xlsx)
        hang(t1, 5, "Thư mục kết quả", v_ra, filedialog.askdirectory)
        v_lech, v_ktt = hai_hop_kiem(t1, 6)
        pb1, lb1, tt1 = khung(t1, 9)

        def doi_chieu_2tab():
            xl, ra, kho = v_xl.get(), v_ra.get() or TM["2_doi_chieu"], v_kho.get().strip()
            if not os.path.isfile(xl): raise ValueError("Chưa chọn file Excel tổng hợp")
            if kho and not os.path.isfile(kho): raise ValueError(f"Không tìm thấy file kho có sẵn: {kho}")
            if not kho and not os.path.isdir(v_tm.get()): raise ValueError("Chưa chọn thư mục kho (hoặc file kho có sẵn)")
            log, (bao, het), T0 = logger(tt1), tien_do(pb1, lb1), time.time()
            buoc = ([] if kho else ["Quét kho"]) + ["Khớp tên sách"] + [b for b, chon in (
                (NHAN_LECH, v_lech.get()), (NHAN_KTT, v_ktt.get())) if chon]
            ten = lambda b: f"Bước {buoc.index(b) + 1}/{len(buoc)} - {b}" if b in buoc else ""
            doc_excel(xl)  # sai file Excel thì báo ngay, trước khi quét kho lâu
            if not kho:
                kho = buoc_quet(v_tm.get(), v_ext.get().split(), log, bao, ten("Quét kho"))
                log("  (lần sau muốn bỏ qua bước quét: chọn file kho này ở ô 'File kho có sẵn')")
            kq, files = khop_ten(xl, kho, log, bao, ten("Khớp tên sách"))
            buoc_noi_dung(kq, files, ten(NHAN_LECH), ten(NHAN_KTT), log, bao)
            buoc_ghi(kq, files, xl, ra, log, bao)
            het(f"Xong - tổng thời gian {_gio(time.time() - T0)}")
        b1 = ttk.Button(t1, text="Chạy đối chiếu", command=lambda: chay(b1, bat(pb1, doi_chieu_2tab), tt1))
        b1.grid(row=8, column=1, sticky="w", pady=6)
    else:
        # --- tab 1: quét kho ---
        t1 = ttk.Frame(nb, padding=10); nb.add(t1, text="1. Quét kho")
        v_tm, v_ext = tk.StringVar(), tk.StringVar(value=EXT_MAC_DINH)
        hang(t1, 0, "Thư mục kho", v_tm, filedialog.askdirectory)
        hang(t1, 1, "Định dạng (cách)", v_ext)
        pb1, lb1, tt1 = khung(t1, 3)

        def quet_4tab():
            if not os.path.isdir(v_tm.get()): raise ValueError("Chưa chọn thư mục kho")
            log, (bao, het), T0 = logger(tt1), tien_do(pb1, lb1), time.time()
            kho = buoc_quet(v_tm.get(), v_ext.get().split(), log, bao, "Quét kho")
            ui(v_kho2.set, kho); ui(v_kho3.set, kho)  # tab 2, 3 lấy luôn file kho mới
            het(f"Xong - tổng thời gian {_gio(time.time() - T0)}")
        b1 = ttk.Button(t1, text="Quét", command=lambda: chay(b1, bat(pb1, quet_4tab), tt1)); b1.grid(row=2, column=1, sticky="w", pady=6)

        # --- tab 2: đối chiếu Excel (khớp tên sách) ---
        t2 = ttk.Frame(nb, padding=10); nb.add(t2, text="2. Đối chiếu Excel")
        v_xl2, v_kho2, v_ra2 = tk.StringVar(value=XL_MD), tk.StringVar(value=moi_nhat(*KHO)), tk.StringVar(value=TM["2_doi_chieu"])
        o_excel.append(v_xl2)
        hang(t2, 0, "File Excel tổng hợp", v_xl2, chon_xlsx)
        hang(t2, 1, "File kho (xlsx)", v_kho2, chon_csv)
        hang(t2, 2, "Thư mục kết quả", v_ra2, filedialog.askdirectory)
        pb2, lb2, tt2 = khung(t2, 4)

        def khop_4tab():
            xl, kho = v_xl2.get(), v_kho2.get()
            if not os.path.isfile(xl): raise ValueError("Chưa chọn file Excel tổng hợp")
            if not os.path.isfile(kho): raise ValueError("Chưa chọn file kho")
            log, (bao, het), T0 = logger(tt2), tien_do(pb2, lb2), time.time()
            kq, files = khop_ten(xl, kho, log, bao, "Khớp tên sách")
            ra1 = buoc_ghi(kq, files, xl, v_ra2.get() or TM["2_doi_chieu"], log, bao)
            ui(v_dc3.set, ra1)
            het(f"Xong - tổng thời gian {_gio(time.time() - T0)}")
        b2 = ttk.Button(t2, text="Đối chiếu", command=lambda: chay(b2, bat(pb2, khop_4tab), tt2)); b2.grid(row=3, column=1, sticky="w", pady=6)

        # --- tab 3: kiểm tra nội dung ---
        t3 = ttk.Frame(nb, padding=10); nb.add(t3, text="3. Kiểm tra nội dung")
        v_dc3, v_xl3 = tk.StringVar(value=moi_nhat(*DC)), tk.StringVar(value=XL_MD)
        v_kho3, v_ra3 = tk.StringVar(value=moi_nhat(*KHO)), tk.StringVar(value=TM["2_doi_chieu"])
        o_excel.append(v_xl3)
        hang(t3, 0, "File doi_chieu_*.xlsx", v_dc3, chon_xlsx)
        hang(t3, 1, "File Excel tổng hợp", v_xl3, chon_xlsx)
        hang(t3, 2, "File kho (xlsx)", v_kho3, chon_csv)
        hang(t3, 3, "Thư mục kết quả", v_ra3, filedialog.askdirectory)
        v_lech, v_ktt = hai_hop_kiem(t3, 4)
        pb3, lb3, tt3 = khung(t3, 7)

        def kiem_4tab():
            dc, xl, kho = v_dc3.get(), v_xl3.get() or v_xl2.get(), v_kho3.get() or v_kho2.get()
            for p, m in ((dc, "file doi_chieu"), (xl, "file Excel tổng hợp"), (kho, "file kho")):
                if not os.path.isfile(p): raise ValueError(f"Chưa chọn {m}")
            log, (bao, het), T0 = logger(tt3), tien_do(pb3, lb3), time.time()
            buoc = [b for b, chon in ((NHAN_LECH, v_lech.get()), (NHAN_KTT, v_ktt.get())) if chon]
            ten = lambda b: f"Bước {buoc.index(b) + 1}/{len(buoc)} - {b}" if b in buoc else ""
            bao("Đọc file đối chiếu và file kho")
            kq, files = doc_doi_chieu(dc), doc_kho(kho)
            log(f"{len(kq)} dòng sách, {len(files)} file kho")
            buoc_noi_dung(kq, files, ten(NHAN_LECH), ten(NHAN_KTT), log, bao)
            buoc_ghi(kq, files, xl, v_ra3.get() or TM["2_doi_chieu"], log, bao)
            het(f"Xong - tổng thời gian {_gio(time.time() - T0)}")
        b3 = ttk.Button(t3, text="Kiểm tra", command=lambda: chay(b3, bat(pb3, kiem_4tab), tt3)); b3.grid(row=6, column=1, sticky="w", pady=6)

    # --- tab cuối: đổi tên & lưu (giống nhau ở cả hai bản) ---
    t4 = ttk.Frame(nb, padding=10); nb.add(t4, text=f"{so_tab}. Đổi tên & lưu")
    v_dc4, v_xl4, v_move = tk.StringVar(value=moi_nhat(*DC)), tk.StringVar(value=XL_MD), tk.BooleanVar()
    v_dich, v_ra4 = tk.StringVar(value=TM["sach_da_luu"]), tk.StringVar(value=TM["3_doi_ten"])
    hang(t4, 0, "File doi_chieu_*.xlsx", v_dc4, chon_xlsx)
    hang(t4, 1, "File Excel tổng hợp", v_xl4, chon_xlsx)
    hang(t4, 2, "Thư mục đích", v_dich, filedialog.askdirectory)
    hang(t4, 3, "Thư mục kết quả", v_ra4, filedialog.askdirectory)
    ttk.Checkbutton(t4, text="Di chuyển file (mặc định: copy)", variable=v_move).grid(row=4, column=0, columnspan=3, sticky="w")
    pb4, lb4, tt4 = khung(t4, 6)
    tt4.config(wrap="none")
    dl = lambda b: f"{b / 2**20:.0f} MB" if b < 2**30 else f"{b / 2**30:.2f} GB"

    def luu(thu):
        xl = v_xl4.get() or next((v.get() for v in o_excel if v.get()), "")
        ra = v_ra4.get() or TM["3_doi_ten"]
        if not os.path.isfile(v_dc4.get()): raise ValueError("Chưa chọn file doi_chieu")
        if not os.path.isdir(v_dich.get()): raise ValueError("Chưa chọn thư mục đích")
        if not thu and not os.path.isfile(xl): raise ValueError("Chưa chọn file Excel tổng hợp")
        log, (bao, het), T0 = logger(tt4), tien_do(pb4, lb4), time.time()
        bao("Đọc file đối chiếu, lập danh sách file cần chép")
        kq = doc_doi_chieu(v_dc4.get()); kh = ke_hoach_luu(kq, v_dich.get())
        so_co = sum(da_co(s, p) for _, s, p in kh)
        so_bia = sum(os.path.basename(os.path.dirname(p)) == "Bia" for _, _, p in kh)
        log(f"{len(kh)} file sang {v_dich.get()} ({so_bia} file bìa vào thư mục Bia): "
            f"{len(kh) - so_co} cần {'di chuyển' if v_move.get() else 'chép'}, {so_co} đã có sẵn")
        if thu:
            log("\n".join(f"{s}\n   → {p}{'   (đã có - bỏ qua)' if da_co(s, p) else ''}" for _, s, p in kh)
                + "\n(chạy thử - chưa copy)")
            het(f"Chạy thử xong - {len(kh)} file, chưa chép file nào")
            return
        buoc = "Di chuyển file" if v_move.get() else "Chép file"
        def prog(i, xb, tb):
            hien = f"{i}/{len(kh)} file, {dl(xb)}/{dl(tb)}"
            bao(buoc, xb, tb, hien) if tb else bao(buoc, i, len(kh), hien)  # tiến độ, thời gian còn lại theo dung lượng
            if i % 100 == 0:
                log(f"  {hien}")
        chep, co_san, loi = luu_file(kh, v_move.get(), prog)
        if DUNG.is_set():
            log(f"ĐÃ DỪNG theo yêu cầu sau {chep + co_san + len(loi)}/{len(kh)} file - chạy lại sẽ bỏ qua file đã chép, làm tiếp phần còn lại.")
        bao("Ghi file kết quả")
        ra1, ra2 = ghi_ket_qua(kq, xl, ra, doc_thu_muc(v_dc4.get()), doc_can_ocr(v_dc4.get()))
        log(f"Xong: chép {chep} file, {co_san} file đã có sẵn (bỏ qua), {len(loi)} lỗi."
            + "".join(f"\n  LỖI {x}" for x in loi[:50]) + f"\nKết quả: {ra1}\nExcel đánh dấu: {ra2}")
        het(f"Xong - {'di chuyển' if v_move.get() else 'chép'} {chep} file, {co_san} file đã có sẵn, {len(loi)} lỗi - tổng thời gian {_gio(time.time() - T0)}")
    b4a = ttk.Button(t4, text="Chạy thử", command=lambda: chay(b4a, bat(pb4, lambda: luu(True)), tt4)); b4a.grid(row=5, column=1, sticky="w", pady=6)
    b4b = ttk.Button(t4, text="Thực hiện", command=lambda: chay(b4b, bat(pb4, lambda: luu(False)), tt4)); b4b.grid(row=5, column=1, sticky="e", pady=6)

    root.mainloop()


# ---------- tự kiểm tra ----------
def _test():
    import tempfile, fitz, docx, openpyxl
    assert chuan_hoa("Lịch sử Chính phủ Việt Nam- T1") == "lich su chinh phu viet nam t1"
    assert chuan_hoa("Đổi mới  công tác\n") == "doi moi cong tac"
    assert ma_ngay("01-2015") == "0115" and ma_ngay("1/2025") == "0125"
    assert ma_ngay("19/1/2007") == "190107" and ma_ngay("11-2-18") == "110218"
    assert ma_ngay(dt.datetime(2008, 1, 8)) == "080108" and ma_ngay(None) == ""
    assert ma_lan_xb(None) == "M00" and ma_lan_xb(1) == "M00" and ma_lan_xb("TB") == "T01"
    assert ma_lan_xb(2) == "T02" and ma_lan_xb("Lần 3") == "T03" and ma_lan_xb("TB lần 2") == "T02"
    assert ten_file_moi("Tên sách", "1/2021", "PL", "TB") == "ten_sach_0121_R_PL_T01"
    assert ten_file_moi("Tên sách", None, "Đ", None) == "ten_sach_R_D_M00"
    # mã BBT quy định (BBT_distinct.xlsx cột F): theo đúng giá trị cột J, ô trống = '(trống)', chưa điền thì giữ BBT bỏ dấu
    wbm = openpyxl.Workbook(); wsm = wbm.active
    for r in (["STT", "BBT", "Mã", "Số", "Năm", "Mã quy định"], [1, "Đảng", "DANG", 1, "", "Đ"], [2, "(trống)", "XX", 1, "", "kxd"],
              [3, "PL", "PL", 1, "", None], [4, "Vụ KH", "VUKH", 1, "", "vụ-kh 2"], [5, 0, "0", 1, "", 0], [6, "KT", "KT", 1, "", 5]):
        wsm.append(r)
    fm = os.path.join(tempfile.gettempdir(), "sxb_test_ma_bbt.xlsx"); wbm.save(fm)
    assert doc_ma_bbt(fm) == {"Đảng": "D", "(trống)": "KXD", "Vụ KH": "VUKH2", "0": "00", "KT": "05"}
    MA_BBT.update(doc_ma_bbt(fm)); os.remove(fm)
    try:
        assert [ten_file_moi("Sách", None, b, None) for b in ("Đảng", None, "PL", " Vụ KH ")] == \
            ["sach_R_D_M00", "sach_R_KXD_M00", "sach_R_PL_M00", "sach_R_VUKH2_M00"]
    finally:
        MA_BBT.clear()
    dai = ten_file_moi("Cuốn sách " + "rất dài " * 40, "01/2021", "PL", "TB")
    assert dai.endswith("_0121_R_PL_T01") and len(dai) <= DAI_TEN + 15 and "__" not in dai and dai.startswith("cuon_sach_rat_dai")
    assert thang_qdxb("05/2023") == "05" and thang_qdxb("19/01/2007") == "01" and thang_qdxb(None) == "00"
    # quét -> xlsx có định dạng -> đọc lại ra dict nội bộ
    with tempfile.TemporaryDirectory() as td:
        p = os.path.join(td, "Sach mau.PDF"); open(p, "wb").write(b"x" * 2048)
        t = dt.datetime(2022, 3, 5, 14, 30).timestamp(); os.utime(p, (t, t))
        kho, bc = os.path.join(td, "kho.xlsx"), {}
        assert quet(td, [".pdf"], kho, bao_cao=bc) == 1 and bc["dinh_dang"] == {"PDF": 1} and bc["loi"] == []
        bc2 = {}; assert quet(os.path.join(td, "khong_co"), [".pdf"], os.path.join(td, "k2.xlsx"), bao_cao=bc2) == 0
        assert len(bc2["loi"]) == 1 and "khong_co" in bc2["loi"][0] and not bc2["loi"][0].startswith("\\\\?")
        DUNG.set()
        try:
            quet(td, [".pdf"], os.path.join(td, "k3.xlsx")); assert False
        except DaDung:
            from openpyxl.worksheet._writer import ALL_TEMP_FILES
            assert not os.path.exists(os.path.join(td, "k3.xlsx")) and not ALL_TEMP_FILES
        finally:
            DUNG.clear()
        ws = openpyxl.load_workbook(kho)["Kho"]
        assert [c.value for c in ws[1]] == COT_KHO and ws["A1"].font.bold and ws.freeze_panes == "A2"
        assert ws.column_dimensions["B"].width == 70 and ws["B2"].alignment.wrap_text and ws["G2"].alignment.wrap_text
        assert [c.value for c in ws[2]] == [1, "Sach mau", "PDF", dt.datetime(2022, 3, 5, 14, 30), 2022, 2, p]
        f, = doc_kho(kho)
        assert (f["dinh_dang"], f["date_modified"], f["nam"], f["path"]) == ("pdf", "2022-03-05 14:30:00", 2022, p)
        cu = os.path.join(td, "cu.csv")  # csv cũ utf-8 phẩy vẫn đọc được
        open(cu, "w", encoding="utf-8-sig", newline="").write("ten_file,dinh_dang,date_modified,path\r\nx,docx,2020-01-02 03:04:05,x.docx\r\n")
        assert doc_kho(cu)[0]["nam"] == 2020
    sach = [dict(nam=2025, dong=4, tt=1, ten="Dạ thưa, xứ Huế bây chừ", lan=None, bbt="QT", qdxb="1/2025"),
            dict(nam=2025, dong=5, tt=2, ten="Luật đầu tư", lan="TB", bbt="PL", qdxb="2/2025"),
            dict(nam=2025, dong=6, tt=3, ten="Không có", lan=None, bbt="PL", qdxb=None)]
    files = [dict(ten_file="Da thua xu Hue bay chu", dinh_dang="docx", date_modified="2025-03-01 10:00:00", path="a.docx"),
             dict(ten_file="Da thua xu Hue bay chu", dinh_dang="docx", date_modified="2025-05-01 10:00:00", path="b.docx"),
             dict(ten_file="Da thua xu Hue bay chu", dinh_dang="pdf", date_modified="2024-05-01 10:00:00", path="c.pdf"),
             dict(ten_file="Luat dau tu (ban in)", dinh_dang="pdf", date_modified="2023-01-01 00:00:00", path="d.pdf"),
             dict(ten_file="Da thua xu Hue bay chu - ban in", dinh_dang="pdf", date_modified="2025-02-01 00:00:00", path="e.pdf")]
    for f in files:
        f["key"] = chuan_hoa(f["ten_file"]); f["nam"] = int(f["date_modified"][:4])
    kq = doi_chieu(sach, files)
    assert kq[0]["tinh_trang"] == "KHỚP" and kq[0]["file_word"] == "b.docx" and kq[0]["file_pdf"] == "e.pdf"
    assert kq[0]["so_file"] == 4 and kq[0]["ghi_chu"] == "1 file khác năm bị bỏ qua"
    assert kq[1]["tinh_trang"].startswith("LỆCH NĂM") and kq[1]["kieu_khop"] == "chứa tên"
    assert kq[2]["tinh_trang"] == "KHÔNG TÌM THẤY FILE"

    # --- quét tiếp sau khi tool tắt giữa chừng: thư mục đã ghi xong không quét lại, thư mục đang dở quét lại ---
    with tempfile.TemporaryDirectory() as td:
        k = os.path.join(td, "kho"); a1, a2 = os.path.join(k, "a"), os.path.join(k, "b")
        for d in (a1, a2):
            os.makedirs(d); open(os.path.join(d, "x.pdf"), "w").close(); open(os.path.join(d, "y.pdf"), "w").close()
        nk, ra = os.path.join(td, "nk.csv"), os.path.join(td, "kho.xlsx")
        with open(nk, "w", encoding="utf-8", newline="") as f:  # giả lập: thư mục a xong (kích thước giả 7777), thư mục b dở
            w = csv.writer(f); w.writerow(["#", os.path.abspath(k), ".pdf"])
            w.writerow(["F", "x", "PDF", "1600000000", "7777", os.path.join(a1, "x.pdf")])
            w.writerow(["F", "y", "PDF", "1600000000", "7777", os.path.join(a1, "y.pdf")]); w.writerow(["D", a1])
            w.writerow(["F", "x", "PDF", "1600000000", "7777", os.path.join(a2, "x.pdf")]); f.write('F,y,PDF,16000')  # dòng ghi dở
        bc = {}
        assert quet(k, [".pdf"], ra, bao_cao=bc, nhat_ky=nk) == 4 and bc["tiep"] == 2 and not os.path.exists(nk)
        fs = doc_kho(ra)
        assert sorted(os.path.basename(os.path.dirname(f["path"])) for f in fs) == ["a", "a", "b", "b"]
        assert openpyxl.load_workbook(ra)["Kho"]["F2"].value == 8  # thư mục a lấy từ nhật ký (7777 byte), không quét lại
        nk2 = os.path.join(td, "nk2.csv")
        with open(nk2, "w", encoding="utf-8", newline="") as f:
            csv.writer(f).writerow(["#", "C:\\khac", ".pdf"])
        try:
            quet(k, [".pdf"], ra, nhat_ky=nk2); assert False
        except ValueError as e:
            assert "thư mục/định dạng khác" in str(e)
        # cache: dấu 'đang đọc' của tiến trình đã chết = file làm tool tắt -> lần sau bỏ qua
        cp = os.path.join(td, "c.sqlite"); c = doc_cache(cp)
        c.db.execute("insert into dang_doc values (?, ?, ?)", (4000000, "k1", "C:\\hong.pdf"))  # pid không tồn tại
        c.db.execute("insert into dang_doc values (?, ?, ?)", (os.getpid(), "k2", "C:\\dang.pdf")); c.dong()  # chính mình
        c = doc_cache(cp)
        assert c.vua_loi == ["C:\\hong.pdf"] and c["k1"] is None and "k2" not in c
        assert c.db.execute("select path from file_loi").fetchall() == [("C:\\hong.pdf",)]
        assert c.doc("k3", "C:\\tot.pdf", lambda: {"nam": []}) == {"nam": []}
        assert c.db.execute("select count(*) from dang_doc where pid = ?", (os.getpid(),)).fetchone()[0] == 0
        c.db.execute("delete from dang_doc"); c.dong()
        c = doc_cache(cp); assert c.vua_loi == []; c.dong()

    # --- ô Excel 32.767 ký tự, thứ tự danh sách file, định dạng thời gian ---
    dai = _ds_file((f"2021-01-01 00:00:00", f"C:\\kho\\{i:06d}" + (".jpg" if i % 2 else ".pdf")) for i in range(2000))
    o = _cell(dai)
    assert len(dai) > 40000 and len(o) <= 32767 and o.endswith("(Excel giới hạn 32.767 ký tự/ô)")
    assert all(p.endswith(".pdf") for _, p in _dong_file(o)[:1000]) and len(_dong_file(o)) == o.count("\n")  # PDF trước, dòng ghi chú bị bỏ
    assert _cell("ngắn") == "ngắn" and _gio(30) == "30 giây" and _gio(600) == "10 phút" and _gio(9000) == "2,5 giờ"

    # --- file kết quả mới nhất ---
    with tempfile.TemporaryDirectory() as td:
        assert moi_nhat(os.path.join(td, "doi_chieu_*.xlsx")) == ""
        for i, ten in enumerate(("doi_chieu_2.xlsx", "doi_chieu_1.xlsx", "khac.xlsx")):
            f = os.path.join(td, ten); open(f, "w").close(); os.utime(f, (1000 + i, 1000 + i))
        assert os.path.basename(moi_nhat(os.path.join(td, "doi_chieu_*.xlsx"), os.path.join(td, "x_*"))) == "doi_chieu_1.xlsx"

    # --- chọn nhầm file: báo lỗi rõ ---
    with tempfile.TemporaryDirectory() as td:
        sai = os.path.join(td, "sai.xlsx"); wb = openpyxl.Workbook(); wb.active.append(["TT", "Tên sách"]); wb.save(sai)
        for ham, chu in ((doc_excel, "không phải file Excel tổng hợp"), (doc_kho, "không phải file kho")):
            try: ham(sai); assert False
            except ValueError as e: assert chu in str(e)

    # --- kho xếp mỗi sách một thư mục ---
    tm = lambda *x: dict(zip(("ten_file", "dinh_dang", "date_modified", "path", "thu_muc"), x))
    kho = [tm("BAN THAO (Thu Hoa)", "doc", "2021-03-01 00:00:00", r"K\Đèn xanh, đèn đỏ\BAN THAO (Thu Hoa).doc", "Đèn xanh, đèn đỏ"),
           tm("DEN XANH (TB)", "pdf", "2021-03-01 00:00:00", r"K\Đèn xanh, đèn đỏ\DEN XANH (TB).pdf", "Đèn xanh, đèn đỏ"),
           tm("Bia gia den xanh", "pdf", "2021-05-01 00:00:00", r"K\Đèn xanh, đèn đỏ\Bia gia den xanh.pdf", "Đèn xanh, đèn đỏ"),
           tm("CHAY (Tr2, 4)", "pdf", "2021-05-01 00:00:00", r"K\Đèn xanh, đèn đỏ\in\CHAY (Tr2, 4).pdf", r"Đèn xanh, đèn đỏ\in"),
           tm("x", "pdf", "2020-01-01 00:00:00", r"K\Đèn xanh, đèn đỏ (1)\x.pdf", "Đèn xanh, đèn đỏ (1)"),
           tm("doi ngoai", "pdf", "2021-01-01 00:00:00", r"K\Đối ngoại quốc phòng\doi ngoai.pdf", "Đối ngoại quốc phòng")]
    kho = [dict(k, path=k["path"].replace("\\", os.sep), thu_muc=k["thu_muc"].replace("\\", os.sep)) for k in kho]  # Linux: /
    import io; buf = io.StringIO(); w = csv.DictWriter(buf, kho[0].keys()); w.writeheader(); w.writerows(kho)
    cp = os.path.join(tempfile.mkdtemp(), "k.csv"); open(cp, "w", encoding="utf-8-sig", newline="").write(buf.getvalue())
    files = doc_kho(cp)
    assert key_thu_muc(os.path.dirname(files[3]["path"])) == ["k", "den xanh den do", "in"]
    assert key_thu_muc(os.path.dirname(files[4]["path"])) == ["k", "den xanh den do"] and "ten_file" not in files[0]
    kq = doi_chieu([dict(nam=2021, dong=2, tt=1, ten="Đèn xanh, đèn đỏ", lan="TB", bbt="QT", qdxb="03/2021"),
                    dict(nam=2021, dong=3, tt=2, ten="Đối ngoại quốc phòng Việt Nam trong thời kỳ mới", lan=None, bbt="QT", qdxb=None),
                    dict(nam=2021, dong=4, tt=3, ten="Đèn xanh, đè", lan=None, bbt="QT", qdxb=None)], files)
    assert kq[0]["kieu_khop"] == "chính xác" and kq[0]["so_file"] == 5 and kq[0]["tinh_trang"] == "KHỚP"
    assert kq[0]["file_word"].endswith(".doc") and kq[0]["file_pdf"].endswith("DEN XANH (TB).pdf")
    assert kq[0]["ghi_chu"] == "1 file khác năm bị bỏ qua; bỏ 2 file bìa/tờ che"
    assert kq[1]["kieu_khop"] == "thư mục rút gọn" and kq[1]["file_pdf"].endswith("doi ngoai.pdf")
    kq2 = doi_chieu([dict(nam=2021, dong=2, tt=1, ten="Đối ngoại quốc phòng", lan=None, bbt="QT", qdxb=None),
                     dict(nam=2021, dong=3, tt=2, ten="Đối ngoại quốc phòng Việt Nam trong thời kỳ mới", lan=None, bbt="QT", qdxb=None),
                     dict(nam=2021, dong=4, tt=3, ten="Đối ngoại quốc phòng (Sách chuyên khảo)", lan=None, bbt="QT", qdxb=None)], files)
    assert kq2[0]["kieu_khop"] == "chính xác" and kq2[1]["so_file"] == 0  # đã có sách trùng tên thư mục -> không khớp rút gọn
    assert kq2[2]["kieu_khop"] == "chính xác" and kq2[2]["tinh_trang"] == "KHỚP"  # bỏ chú thích 'Sách chuyên khảo' khi khớp
    # chú thích sau tên sách không chặn khớp tên -> file Date Modified 2018 khớp ngay dòng 2018, dòng 2009 lệch năm
    assert loi_ten("Ngoại giao và công tác ngoại giao Sách chuyên khảo (Xuất bản lần thứ hai, có sửa chữa)") == "ngoai giao va cong tac ngoai giao"
    assert loi_ten("Chơi đến cùng (Sách tham khảo) Người dịch: Đinh Trọng Minh") == "choi den cung"
    assert loi_ten("Ngoại giao và công tác ngoại giao (Tiếng Lào)").endswith("tieng lao") and loi_ten("Tái bản") == "tai ban"
    ng = [dict(nam=2009, dong=5, tt=1, ten="Ngoại giao và công tác ngoại giao", lan=None, bbt="QT", qdxb="06/2009"),
          dict(nam=2018, dong=6, tt=1, ten="Ngoại giao và công tác ngoại giao Sách chuyên khảo (Xuất bản lần thứ hai)", lan="2", bbt="QT", qdxb="10/2018")]
    fng = [dict(key=chuan_hoa("Ngoại giao và công tác ngoại giao"), dinh_dang="pdf", date_modified="2018-10-20 09:00:00",
                path=r"K\ng\Ngoại giao và công tác ngoại giao.pdf", nam=2018)]
    kng = doi_chieu(ng, fng)
    assert [r["tinh_trang"][:8] for r in kng] == ["LỆCH NĂM", "KHỚP"] and file_theo_dm(kng) == {fng[0]["path"]: 2018}
    assert kq[2]["tinh_trang"] == "KHÔNG TÌM THẤY FILE"  # "den xanh de" chỉ khớp ký tự, không khớp nguyên từ
    assert nam_trang_cuoi("nghi quyet so 41 2017 qh14 bo luat so 100 2015 qh13 in xong va nop luu chieu thang nam 2021") == [2021]
    assert nam_trang_cuoi("in xong va noap lou chieau thaung 3 naem 2021 chiu") == [2021]
    assert nam_trang_cuoi("ky 2021 2026 dang ky xuat ban 538 2021 qd ngay 19 02 2021 nop luu chieu thang 02 nam 2020") == [2020, 2021]
    assert nam_trang_cuoi("asean 2025 so dang ky xuat ban 551 2020 cxbiph ngay 09 9 2020 ma so isbn") == [2020]
    assert nam_trang_cuoi("phu luc 1 my 1900 2020 483 phu luc 2 eu 1900 2020 518") == []  # mục lục, không có mốc XB
    assert nam_trang_dau("dang cam tu h chinh tri quoc gia 2021 280tr 21cm nam 2015 luat so 100 2015") == [2021]
    assert nam_trang_dau("nha xuat ban chinh tri quoc gia su that ha noi 2021 loi nha xuat ban nam 2020 la nam") == [2021]
    assert nam_trang_dau("gioi thieu cuon sach cung ban doc thang 4 nam 2021 nhj xuat ban nam 2004 2012") == [2021]
    assert nam_trang_dau("so dang ky ke hoach xuat ban 4139 2020 cxbiph nop luu chieu thang 10 nam 2020" + " x" * 200 + " 1900 2019") == [2020]
    assert nam_trang_dau("nghi quyet so 133 2020 qh14 ngay 17 11 2020") == []
    assert nam_trang_dau_bac("gioi thieu cung ban doc thang 6 nam 2009 nha xuat ban") == ([2009], 3)
    assert khop_nam(dict(nam=[2018], dau=[2009], dau_bac=3), 2009) == ""  # tái bản giữ lời NXB 2009, trang cuối ghi 2018
    assert khop_nam(dict(nam=[], dau=[2019], dau_bac=3), 2019) == "8 trang đầu"  # trang cuối không có năm -> vẫn dùng
    assert khop_nam(dict(nam=[2020], dau=[2021], dau_bac=2), 2021) == "8 trang đầu"  # CIP đầu sách vẫn thắng
    assert nam_trang_dau("van kien nxb chinh tri quoc gia ha noi 2011 t 7 tr 15") == []  # chú thích tài liệu tham khảo
    assert nam_trang_dau("nop luu chieu thang 4 nam 2021 ma isbn 978 bien muc h chinh tri quoc gia 2020 72tr") == [2021]
    assert nam_trang_dau("h chinh tri quoc gia 2021 244tr gioi thieu cung ban doc thang 11 nam 2019 nha xuat ban") == [2021]
    assert la_bia("Bia 1 chay.pdf") and la_bia("To che x.pdf") and la_bia("Tr 2,4 (x).pdf") and not la_bia("Bia da.pdf", "bia da")
    assert not la_bia("Kinh te bien.pdf")

    # --- tab 3: trích text pdf/docx, cache ---
    tmp = tempfile.mkdtemp()
    pdf, dx, doc = (os.path.join(tmp, n) for n in ("a.pdf", "b.docx", "c.doc"))
    d = fitz.open()
    d.new_page().insert_text((50, 72), "Trang bia")
    d.new_page().insert_text((50, 72), "LOI NOI DAU\nCuon sach Luat dau tu nay gioi thieu...")
    d.new_page().insert_text((50, 72), "NHA XUAT BAN CHINH TRI QUOC GIA\nIn xong va nop luu chieu thang 5 nam 2023")
    d.save(pdf); d.close()
    w = docx.Document()
    for s in ("Bia", "Lời Nhà xuất bản", "Cuốn sách Luật đầu tư được tái bản.", "Nội dung", "Nộp lưu chiểu tháng 6 năm 2019"):
        w.add_paragraph(s)
    w.save(dx)
    open(doc, "wb").write(b"\xd0\xcf\x11\xe0 junk")
    t = trich(pdf)
    assert t["nam"] == [2023] and t["dau"] == [2023] and t["lnd"].startswith("cuon sach luat dau tu")
    assert trich(dx)["nam"] == [2019] and "luat dau tu" in trich(dx)["lnd"] and trich(doc) is None
    assert _docx_text(dx) == "Bia\nLời Nhà xuất bản\nCuốn sách Luật đầu tư được tái bản.\nNội dung\nNộp lưu chiểu tháng 6 năm 2019"
    try:  # .doc thật: tạo bằng Word từ b.docx rồi đọc lại (bỏ qua nếu máy không có Word)
        w = os.name == "nt" and _word_app()
        if w:
            doc2 = os.path.join(tmp, "e.doc"); dd = w.Documents.Open(dx); dd.SaveAs2(doc2, FileFormat=0); dd.Close(False)
            t2 = trich(doc2)
            assert t2["nam"] == [2019] and t2["dau"] == [2019] and "luat dau tu" in t2["lnd"], t2
            dai = os.path.join(tmp, ("thu muc ten sach rat dai " * 6).strip(), ("ban thao " * 10).strip() + ".doc")  # > 255 ký tự
            os.makedirs(duong_dan_dai(os.path.dirname(dai))); shutil.copy2(doc2, duong_dan_dai(dai))
            assert len(dai) > 255 and trich(dai) == t2 and not glob.glob(os.path.join(tempfile.gettempdir(), "sxb_*.doc"))
            import time
            mod, pid, goi = sys.modules[__name__], _WORD[1], []
            assert pid in _pid_word()
            # hẹn giờ: hết giờ thì gọi hàm tắt Word và lần sau mở Word mới (thay hàm tắt bằng hàm ghi nhận để kiểm tra)
            cu = mod.WORD_HET_GIO, mod._giet_word
            mod.WORD_HET_GIO, mod._giet_word = 0.001, lambda: goi.append(1)
            try:
                trich(doc2)
                assert goi and _WORD[0] is None
            finally:
                mod.WORD_HET_GIO, mod._giet_word = cu
                _WORD[1] = pid; _giet_word(); _WORD[1] = None
            # Word bị tắt thật giữa chừng -> file đó bỏ qua, lần sau tự mở Word mới
            assert trich(doc2)["nam"] == [2019]
            pid = _WORD[1]; _giet_word()
            for _ in range(50):
                if pid not in _pid_word(): break
                time.sleep(0.1)
            assert trich(doc2) is None and _WORD[0] is None
            assert trich(doc2)["nam"] == [2019] and _WORD[1] not in (None, pid)
    finally:
        dong_word()
    assert nam_ten_file("x/CHU NGHIA TU BAN 2021.pdf", "chu nghia tu ban 1900 2020") == {2021}
    assert nam_ten_file("x/BAO CAO NAM 2020 (tb2021).pdf", "bao cao nam 2020") == {2021} and nam_ten_file("x/a_090321R08.pdf") == set()
    cache = {}
    assert trich_cache(cache, pdf) == t and len(cache) == 1 and trich_cache(cache, "khong_co.pdf") is None
    cj = os.path.join(tmp, "c.sqlite"); ghi_cache(cache, cj); c = doc_cache(cj)
    assert len(c) == 1 and c[next(iter(cache))] == t and "x" not in c and c.get("x", 0) == 0
    c["moi"] = None; ghi_cache(c, cj); c.dong(); c = doc_cache(cj); assert len(c) == 2 and c["moi"] is None; c.dong()
    assert loi_noi_dau("Mở đầu. LỜI NHÀ XUẤT BẢN Năm 2020 ... xuất bản cuốn sách Đèn xanh, đèn đỏ. Trân trọng") == \
        "cuon sach den xanh den do tran trong" and loi_noi_dau("Không có lời nói đầu") == ""

    bt = os.path.join(tmp, "ban thao 2021.docx"); w = docx.Document(); w.add_paragraph("Ban thao"); w.save(bt)  # không có năm trong nội dung
    r = dict(nam=2023, ten="Luật đầu tư", tinh_trang="LỆCH NĂM - cần mở file kiểm tra", file_word="", file_pdf="", ghi_chu="",
             cac_file=f"2021-01-01 00:00:00  {pdf}\n2021-01-02 00:00:00  {doc}\n2021-01-03 00:00:00  {dx}\n2021-01-04 00:00:00  {os.path.join(tmp, 'bia.jpg')}\n2021-01-05 00:00:00  {bt}")
    xu_ly_lech_nam(r, cache)
    assert r["tinh_trang"] == "KHỚP" and r["file_pdf"] == pdf and r["ghi_chu"] == "khớp theo nội dung (trang cuối)"
    assert r["file_word"] == doc  # .doc cùng thư mục được nhận; .docx ghi năm 2019 hay tên file ghi 2021 thì không
    # trang cuối ghi lưu chiểu 2020, CIP ở đầu sách ghi 2021 -> dòng 2020 khớp trang cuối, dòng 2021 khớp 8 trang đầu
    hai = os.path.join(tmp, "hai", "h.pdf"); os.makedirs(os.path.dirname(hai)); d = fitz.open()
    d.new_page().insert_text((50, 72), "H. : Chinh tri quoc gia, 2021. - 280tr.")
    for _ in range(9): d.new_page().insert_text((50, 72), "Noi dung nam 2015")
    d.new_page().insert_text((50, 72), "In xong va nop luu chieu thang 9 nam 2020"); d.save(hai); d.close()
    assert trich(hai)["nam"] == [2020] and trich(hai)["dau"] == [2021]
    for nam, nguon in ((2020, "trang cuối"), (2021, "8 trang đầu")):
        rh = dict(r, nam=nam, tinh_trang="LỆCH NĂM", cac_file="2021-01-01 00:00:00  " + hai); xu_ly_lech_nam(rh, cache)
        assert rh["tinh_trang"] == "KHỚP" and rh["ghi_chu"] == f"khớp theo nội dung ({nguon})"
    rh = dict(r, nam=2015, tinh_trang="LỆCH NĂM", cac_file="2021-01-01 00:00:00  " + hai); xu_ly_lech_nam(rh, cache)
    assert rh["tinh_trang"] == "LỆCH NĂM" and "thấy: 2020, 2021" in rh["ghi_chu"]  # năm trong nội dung không tính
    # file đã thuộc sách năm 2020 theo Date Modified -> dòng 2021 không đọc lại nội dung file đó
    rh = dict(r, nam=2021, tinh_trang="LỆCH NĂM", cac_file="2020-06-01 00:00:00  " + hai); xu_ly_lech_nam(rh, cache, {hai: 2020})
    assert rh["tinh_trang"] == "LỆCH NĂM" and "h.pdf đã thuộc sách năm 2020 theo Date Modified" in rh["ghi_chu"]
    rh = dict(r, nam=2015, tinh_trang="LỆCH NĂM", cac_file="2021-01-01 00:00:00  " + hai); xu_ly_lech_nam(rh, cache)
    assert rh["tinh_trang"] == "LỆCH NĂM" and "thấy: 2020, 2021" in rh["ghi_chu"]  # năm trong nội dung không tính
    r2 = dict(r, nam=2020, tinh_trang="LỆCH NĂM - cần mở file kiểm tra"); xu_ly_lech_nam(r2, cache)
    assert r2["tinh_trang"].startswith("LỆCH NĂM") and "thấy: 2019, 2023" in r2["ghi_chu"] and "bia.jpg" not in r2["ghi_chu"]

    kho = [dict(ten_file="a", dinh_dang="pdf", date_modified="2021-01-01 00:00:00", path=pdf, nam=2021),
           dict(ten_file="b", dinh_dang="docx", date_modified="2021-06-01 00:00:00", path=dx, nam=2021),
           dict(ten_file="c", dinh_dang="doc", date_modified="2021-06-01 00:00:00", path=doc, nam=2021)]
    assert len(loc_files(kho)) == 3 and len(loc_files(kho, "", "2020")) == 0 and len(loc_files(kho, tmp, "2020-2021")) == 3
    assert loc_files(kho, os.path.join(tmp, "khac")) == []
    rows = [dict(nam=2023, ten="Luật đầu tư", tinh_trang="KHÔNG TÌM THẤY FILE", file_word="", file_pdf="", ghi_chu="", cac_file=""),
            dict(nam=2023, ten="Không có", tinh_trang="KHÔNG TÌM THẤY FILE", file_word="", file_pdf="", ghi_chu="", cac_file=""),
            dict(nam=2019, ten="Luật đầu tư", tinh_trang="KHÔNG TÌM THẤY FILE", file_word="", file_pdf="", ghi_chu="", cac_file=""),
            dict(nam=2018, ten="Luật tư", tinh_trang="KHÔNG TÌM THẤY FILE", file_word="", file_pdf="", ghi_chu="", cac_file=""),
            dict(nam=2023, ten="Đầu tư", tinh_trang="KHÔNG TÌM THẤY FILE", file_word="", file_pdf="", ghi_chu="", cac_file="")]
    DUNG.set()
    try:
        xu_ly_khong_tim_thay([dict(r) for r in rows], loc_files(kho), cache); assert False
    except DaDung:
        pass
    finally:
        DUNG.clear()
    xu_ly_khong_tim_thay(rows, loc_files(kho), cache)
    assert rows[0]["tinh_trang"] == "KHỚP" and rows[0]["file_pdf"] == pdf and rows[0]["so_file"] == 2 and rows[0]["file_word"] == doc
    assert rows[1]["tinh_trang"] == "KHÔNG TÌM ĐƯỢC"
    assert rows[2]["tinh_trang"] == "KHỚP" and rows[2]["file_word"] == dx and rows[2]["file_pdf"] == ""
    assert rows[3]["tinh_trang"] == "KHÔNG TÌM ĐƯỢC"  # "luat tu" không là cụm liền trong lời nói đầu
    assert rows[4]["tinh_trang"] == "KHÔNG TÌM ĐƯỢC"  # "dau tu" có trong lời nói đầu nhưng không đứng sau "cuốn sách"
    assert nhac_ten_sach("nxb xuat ban cuon sach chuyen khao luat dau tu", "luat dau tu")
    assert not nhac_ten_sach("can cu hien phap nuoc cong hoa nam 2013", "hien phap nuoc cong hoa")

    # --- PDF cần OCR: ảnh scan / chữ dạng hình vẽ -> ghi chú dòng sách, sheet thư mục, sheet CAN_OCR ---
    tm_ocr = os.path.join(tmp, "ocr"); os.makedirs(tm_ocr)
    scan, ve, bia_scan = (os.path.join(tm_ocr, n) for n in ("s.pdf", "v.pdf", "Bia 1.pdf"))
    for f in (scan, bia_scan):
        d = fitz.open(); pix = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 20, 20)); pix.clear_with(200)
        for _ in range(3): d.new_page().insert_image(fitz.Rect(0, 0, 200, 200), pixmap=pix)
        d.save(f); d.close()
    d = fitz.open()
    for _ in range(3): d.new_page().draw_rect(fitz.Rect(10, 10, 100, 100))
    d.save(ve); d.close()
    assert trich(scan)["ocr"] == "scan" and trich(scan)["trang"] == 3 and trich(ve)["ocr"] == "hinh_ve" and trich(pdf)["ocr"] == ""
    ro = dict(nam=2021, dong=9, ten="Sách scan", tinh_trang="LỆCH NĂM", file_word="", file_pdf="", ghi_chu="",
              cac_file="\n".join(f"2021-01-01 00:00:00  {f}" for f in (scan, ve, bia_scan)))
    xu_ly_lech_nam(ro, cache)
    assert ro["tinh_trang"] == "LỆCH NĂM" and ro["ghi_chu"].endswith("cần OCR (PDF không có lớp chữ): s.pdf (3 trang), v.pdf (3 trang)")
    kho_ocr = [dict(dinh_dang="pdf", path=f) for f in (scan, ve, bia_scan, pdf)]
    ocr_ds = ds_can_ocr([ro], kho_ocr, cache)
    assert [(x["ten_file"], x["trang"], x["ly_do"]) for x in ocr_ds] == [("s.pdf", 3, LY_DO_OCR["scan"]), ("v.pdf", 3, LY_DO_OCR["hinh_ve"])]
    assert ocr_ds[0]["sach"] == "2021 - dòng 9 - Sách scan" and ocr_ds[0]["ten_tm"] == "ocr" and ocr_ds[0]["uu_tien"] == OCR_CAO
    khop = dict(ro, tinh_trang="KHỚP", file_pdf=pdf, file_word="")  # thư mục khác đã khớp -> file OCR ở thư mục đó ưu tiên thấp
    assert [x["uu_tien"] for x in ds_can_ocr([ro, dict(khop, file_pdf=scan)], kho_ocr, cache)] == [OCR_THAP, OCR_THAP]
    g = next(x for x in doi_chieu_thu_muc([ro], kho_ocr, cache) if x["ten_tm"] == "ocr")["ghi_chu"]
    assert "2 file PDF cần OCR (không có lớp chữ) - xem sheet CAN_OCR" in g and "không thấy năm" not in g

    # --- ghi / đọc lại doi_chieu, tab 4 ---
    xl = os.path.join(tmp, "th.xlsx"); wb = openpyxl.Workbook(); ws = wb.active; ws.title = "2023"
    ws.append(["TT", "Tên sách"]); ws.append([1, "Luật đầu tư"]); ws.append([2, "Không có"]); wb.save(xl)
    kq = [dict(nam=2023, dong=2, tt=1, ten="Luật đầu tư", lan="TB", bbt="PL", qdxb="05/2023", tinh_trang="KHỚP",
               kieu_khop="chính xác", so_file=2, file_word=dx, file_pdf=pdf, ten_moi="luat_dau_tu_0523_R_PL_T01",
               ghi_chu="", cac_file=""),
          dict(nam=2023, dong=3, tt=2, ten="Không có", lan=None, bbt="PL", qdxb=None, tinh_trang="KHÔNG TÌM THẤY FILE",
               kieu_khop="", so_file=0, file_word="", file_pdf="", ten_moi="khong_co_R_PL_M00", ghi_chu="", cac_file="")]
    kho_tm = [dict(dinh_dang="docx", path=dx), dict(dinh_dang="pdf", path=pdf),
              dict(dinh_dang="jpg", path=os.path.join(tmp, "x", "bia.jpg")), dict(dinh_dang="pdf", path=os.path.join(tmp, "y", "z.pdf"))]
    kq[1]["cac_file"] = "2023-01-01 00:00:00  " + os.path.join(tmp, "y", "z.pdf")
    trung = os.path.join(tmp, "z_trung", "a.pdf"); os.makedirs(os.path.dirname(trung)); shutil.copy2(pdf, trung)
    kho_tm.append(dict(dinh_dang="pdf", path=trung)); kq[0]["cac_file"] = "2023-01-01 00:00:00  " + trung
    assert doi_chieu_thu_muc(kq, kho_tm, cache)[-1]["tinh_trang"] == TM_CHUA  # bản trùng chưa đọc -> chưa biết năm
    trich_cache(cache, trung)
    tm = doi_chieu_thu_muc(kq, kho_tm, cache)
    assert [(x["ten_tm"], x["tinh_trang"], x["so_file"]) for x in tm] == [
        (os.path.basename(tmp), TM_KHOP, 2), ("x", TM_KHONG, 1), ("y", TM_CHUA, 1), ("z_trung", TM_TRUNG, 1)]
    assert tm[3]["nam_cuoi"] == "trang cuối: 2023; 8 trang đầu: 2023"
    assert tm[0]["nam_cuoi"] == "trang cuối: 2019, 2023; 8 trang đầu: 2019, 2023" and tm[0]["ghi_chu"] == "mở được 2/2 file Word/PDF"
    assert tm[1]["ghi_chu"].startswith("không có file Word/PDF") and tm[2]["ghi_chu"].startswith("mở được 0/1")
    assert tm[0]["loai"] == "DOCX 1, PDF 1" and tm[0]["sach_khop"] == "2023 - dòng 2 - Luật đầu tư" and tm[2]["sach_khac"].endswith("Không có")
    ra1, ra2 = ghi_ket_qua(kq, xl, os.path.join(tmp, "ra"), tm, ocr_ds)
    assert doc_thu_muc(ra1) == tm and doc_can_ocr(ra1) == ocr_ds
    kq2 = doc_doi_chieu(ra1)
    assert [r["ten"] for r in kq2] == ["Luật đầu tư", "Không có"] and kq2[0]["file_pdf"] == pdf and kq2[0]["dong"] == 2
    kq2.append(dict(kq2[0], dong=4, qdxb=None))  # trùng tên file mới -> _2
    dich = os.path.join(tmp, "dich"); kh = ke_hoach_luu(kq2, dich)
    t = "luat_dau_tu_0523_R_PL_T01"
    assert [os.path.relpath(p, dich) for _, _, p in kh] == [
        os.path.join("2023", "thang 05", t, t + ".docx"), os.path.join("2023", "thang 05", t, t + ".pdf"),
        os.path.join("2023", "thang 00", t, t + ".docx"), os.path.join("2023", "thang 00", t, t + ".pdf")]
    kq2[2]["qdxb"] = "05/2023"; kh = ke_hoach_luu(kq2, dich)  # trùng thư mục -> thư mục _2
    assert kh[2][2] == os.path.join(dich, "2023", "thang 05", t + "_2", t + "_2.docx") and kh[3][2].endswith(t + "_2.pdf")
    assert luu_file(kh) == (4, 0, []) and all(os.path.isfile(p) for _, _, p in kh) and os.path.isfile(pdf)
    assert kq2[0]["dich"] == kh[0][2] + "\n" + kh[1][2] and not glob.glob(os.path.join(dich, "**", "*.dang_chep"), recursive=True)
    # chạy lại: dùng lại đúng thư mục cũ, không chép lại
    kh2 = ke_hoach_luu([dict(r, dich="") for r in kq2], dich)
    assert [p for _, _, p in kh2] == [p for _, _, p in kh] and luu_file(kh2) == (0, 4, [])
    # thư mục trùng tên nhưng là sách khác (file khác dung lượng) -> sang _3; copy dở (thiếu PDF) -> dùng lại
    kh3 = ke_hoach_luu([dict(kq2[0], file_word="", file_pdf=hai)], dich)
    assert kh3[0][2] == os.path.join(dich, "2023", "thang 05", t + "_3", t + "_3.pdf")
    os.remove(kh[1][2]); kh4 = ke_hoach_luu([dict(kq2[0], dich="")], dich)
    assert [p for _, _, p in kh4] == [p for _, _, p in kh[:2]] and luu_file(kh4) == (1, 1, [])
    # bấm Dừng: không chép thêm file nào
    DUNG.set()
    try:
        assert luu_file(ke_hoach_luu([dict(kq2[0], ten_moi="sach_dung", dich="")], dich)) == (0, 0, [])
        assert not os.path.exists(os.path.join(dich, "2023", "thang 05", "sach_dung"))
    finally:
        DUNG.clear()
    # bìa: ảnh, .ai, PDF 1-2 trang, PDF tên 'Bìa' -> <thư mục sách>/Bia/<tên _B_>; PDF 1-2 trang không làm ruột
    assert ten_bia("luat_dau_tu_0523_R_PL_T01") == "luat_dau_tu_0523_B_PL_T01" and ten_bia("ten_R_R_M00") == "ten_B_R_M00"
    tb = os.path.join(tmp, "co_bia"); os.makedirs(tb)
    def lam_pdf(ten, n):
        d = fitz.open()
        for _ in range(n): d.new_page().insert_text((50, 72), "Noi dung")
        d.save(os.path.join(tb, ten)); d.close(); return os.path.join(tb, ten)
    ruot, mong, bia4 = lam_pdf("ruot.pdf", 5), lam_pdf("mong 1 trang.pdf", 1), lam_pdf("Bia 4.pdf", 3)
    anh, ai_, bt2 = os.path.join(tb, "bia.jpg"), os.path.join(tb, "bia.ai"), os.path.join(tb, "ban thao.docx")
    for f, nd in ((anh, b"jpg"), (ai_, b"ai"), (bt2, b"x")): open(f, "wb").write(nd)
    ds = [("2023-02-01 00:00:00", mong), ("2023-01-01 00:00:00", ruot), ("2023-01-01 00:00:00", bia4), ("2023-01-01 00:00:00", bt2)]
    assert chon_file(ds)[1] == ruot  # PDF 1 trang mới hơn vẫn không làm ruột
    d = fitz.open(); pix = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 400, 400)); pix.clear_with(90)
    for _ in range(3): d.new_page().insert_image(fitz.Rect(0, 0, 400, 400), pixmap=pix)
    ban_in = os.path.join(tb, "ban in.pdf"); d.save(ban_in); d.close()
    assert os.path.getsize(ban_in) > os.path.getsize(ruot)  # nặng hơn nhưng ít trang hơn
    assert chon_file([("2023-01-01 00:00:00", ban_in), ("2023-01-01 00:00:00", ruot)])[1] == ruot  # cùng ngày: nhiều trang hơn
    # trải bìa nằm ngang + 2 tai gập (3 trang) là bìa, không làm ruột
    d = fitz.open(); d.new_page(width=1148, height=567); d.new_page(width=348, height=567); d.new_page(width=348, height=567)
    trai = os.path.join(tb, "trai_sach.pdf"); d.save(trai); d.close()
    assert pdf_mong(trai) and not pdf_mong(ruot)
    assert chon_file([("2023-02-01 00:00:00", trai), ("2023-01-01 00:00:00", ruot)])[1] == ruot
    chi_bia = dict(kq2[0], file_word="", file_pdf="", ten_moi="chi_co_bia_0523_R_PL_T01", dich="",
                   cac_file=f"2023-02-01 00:00:00  {trai}\n2023-01-01 00:00:00  {anh}")
    assert sorted(os.path.basename(p) for _, _, p in ke_hoach_luu([chi_bia], dich)) == ["chi_co_bia_0523_B_PL_T01.jpg", "chi_co_bia_0523_B_PL_T01.pdf"]  # chỉ có bìa: vẫn lưu bìa
    rb = dict(kq2[0], file_word=bt2, file_pdf=ruot, ten_moi="sach_co_bia_0523_R_PL_T01", dich="",
              cac_file="\n".join(f"2023-01-01 00:00:00  {f}" for f in (ruot, mong, bia4, anh, ai_, bt2)))
    assert ds_bia(rb) == sorted([mong, bia4, anh, ai_])
    khb = ke_hoach_luu([rb], dich); tmb = os.path.dirname(khb[0][2])
    assert [os.path.relpath(p, tmb) for _, _, p in khb] == ["sach_co_bia_0523_R_PL_T01.docx", "sach_co_bia_0523_R_PL_T01.pdf",
        os.path.join("Bia", "sach_co_bia_0523_B_PL_T01.pdf"), os.path.join("Bia", "sach_co_bia_0523_B_PL_T01.ai"),
        os.path.join("Bia", "sach_co_bia_0523_B_PL_T01.jpg"), os.path.join("Bia", "sach_co_bia_0523_B_PL_T01_2.pdf")]
    assert luu_file(khb) == (6, 0, []) and luu_file(ke_hoach_luu([dict(rb, dich="")], dich)) == (0, 6, [])  # chạy lại: bỏ qua
    open(os.path.join(tmb, "Bia", "sach_co_bia_0523_B_PL_T01.png"), "wb").write(b"file khac")  # tên đã có file khác
    png = os.path.join(tb, "bia.png"); open(png, "wb").write(b"png moi")
    khp = ke_hoach_luu([dict(rb, dich="", cac_file=rb["cac_file"] + f"\n2023-01-01 00:00:00  {png}")], dich)
    assert next(d for _, s, d in khp if s == png).endswith("sach_co_bia_0523_B_PL_T01_2.png")  # không ghi đè file khác
    # đọc song song: tiến trình con sập chỉ bỏ qua đúng file gây sập, các file khác vẫn đọc được; chạy lại lấy từ cache
    ss = [lam_pdf(f"ss{i}.pdf", 3) for i in range(6)] + [lam_pdf("sap.pdf", 3)]
    cs = Cache(os.path.join(tmp, "ss.sqlite"))
    assert doc_truoc(cs, ss + [anh, ss[0]], 2, ham=_trich_thu) == 7 and cs.vua_loi == [ss[-1]]
    assert all(cs.get(_khoa_cache(p))["trang"] == 3 for p in ss[:-1]) and cs.get(_khoa_cache(ss[-1]), 1) is None
    assert cs.db.execute("select count(*) from file_loi").fetchone()[0] == 1 and doc_truoc(cs, ss, 2) == 0
    cs.dong()
    # lỗi một file không làm dừng cả lượt
    hong = dict(kq2[0], file_word=os.path.join(tmp, "khong_co.docx"), ten_moi="sach_loi", dich="", ghi_chu="")
    goi = []
    chep, co_san, loi = luu_file(ke_hoach_luu([hong, dict(kq2[0], ten_moi="sach_tot", dich="")], dich), progress=lambda *a: goi.append(a))
    assert chep == 3 and len(loi) == 1 and "LỖI lưu khong_co.docx" in hong["ghi_chu"]
    tong = os.path.getsize(pdf) * 2 + os.path.getsize(dx)  # file nguồn không tồn tại không tính vào tổng dung lượng
    assert [a[0] for a in goi] == [1, 2, 3, 4] and goi[-1][1:] == (tong, tong) and goi[0][1] == 0
    ra1, ra2 = ghi_ket_qua(kq2, xl, os.path.join(tmp, "ra"))
    ws = openpyxl.load_workbook(ra2)["2023"]
    assert ws.cell(1, 6).value == "Đường dẫn đích" and ws.cell(2, 6).value == kq2[0]["dich"] and ws.cell(2, 3).value == "KHỚP"
    assert all(ws.cell(2, c).fill.fgColor.rgb == "00FFFF00" for c in range(1, 7)) and ws.cell(3, 1).fill.fill_type is None  # KHỚP tô vàng
    shutil.rmtree(tmp, ignore_errors=True)
    print("OK")


if __name__ == "__main__":
    import multiprocessing; multiprocessing.freeze_support()  # exe: tiến trình con đọc PDF song song
    import sxb_tool  # chạy qua module sxb_tool (không phải __main__): tiến trình con của exe mới gọi được trich()
    sxb_tool._test() if "--test" in sys.argv else sxb_tool.gui(4 if "--4tab" in sys.argv else 2)
