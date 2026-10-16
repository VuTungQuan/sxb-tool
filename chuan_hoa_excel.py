"""Chuẩn hoá 'Tong hop SXB 2006-2025.xlsx': mọi sheet năm cùng một bộ cột, header ở dòng 1.

Chạy: python chuan_hoa_excel.py <excel_goc.xlsx> [excel_ra.xlsx]
"""
import re, sys, datetime as dt
import openpyxl
from openpyxl.styles import Font
from sxb_tool import chuan_hoa

COT_CHUAN = ["TT", "Tên sách", "Tác giả", "Phương thức XB", "Lần XB", "Số lượng in", "Số trang", "Khổ sách",
             "Giá bìa", "BBT", "QĐXB", "Nơi in", "Ghi chú", "ISBN", "Số trang quy đổi"]


def ten_cot(h, idx):
    if h is None:
        return "Lần XB" if idx == 4 else None  # sheet 2014 để trống header cột E
    k = chuan_hoa(h)
    if k == "tt": return "TT"
    if "ten sach" in k: return "Tên sách"
    if "tac gia" in k: return "Tác giả"
    if "phuong thuc" in k: return "Phương thức XB"
    if k.startswith("lan x"): return "Lần XB"
    if "so luong" in k or k == "sl in": return "Số lượng in"
    if "quy doi" in k and "tong" not in k: return "Số trang quy đổi"
    if k in ("trang", "so trang"): return "Số trang"
    if k.startswith("kho"): return "Khổ sách"
    if k.startswith("gia bia"): return "Giá bìa"
    if k in ("bbt", "ban bt", "ban bien tap"): return "BBT"
    if "qdxb" in k or "thoi gian" in k: return "QĐXB"
    if "noi in" in k: return "Nơi in"
    if "ghi chu" in k: return "Ghi chú"
    if "isbn" in k: return "ISBN"
    return re.sub(r"\s+", " ", str(h)).strip()  # cột lạ: giữ nguyên tên, xếp cuối


def sach_text(v):
    return re.sub(r"\s+", " ", str(v)).strip() if isinstance(v, str) else v


def chuan_qdxb(v):
    if isinstance(v, (dt.datetime, dt.date)):
        return v.strftime("%d/%m/%Y")
    if v is None:
        return None
    nums = re.findall(r"\d+", str(v))
    if len(nums) == 3:
        d, m, y = nums
        return f"{int(d):02d}/{int(m):02d}/{y if len(y) == 4 else '20' + y[-2:]}"
    if len(nums) == 2:
        m, y = nums
        return f"{int(m):02d}/{y if len(y) == 4 else '20' + y[-2:]}"
    return sach_text(v)


def chuan_lan_xb(v):
    v = sach_text(v)
    return None if v in ("", "`") else v


def chuan_hoa_excel(goc, ra):
    wb = openpyxl.load_workbook(goc, read_only=True, data_only=True)
    out = openpyxl.Workbook(); del out["Sheet"]
    bao_cao = []
    for sh in wb.sheetnames:
        if not sh.strip().isdigit():
            ws_o = out.create_sheet(sh)
            for row in wb[sh].iter_rows(values_only=True): ws_o.append(list(row))
            continue
        hdr, rows, bo = None, [], 0
        for row in wb[sh].iter_rows(values_only=True):
            if hdr is None:
                if any(c and "ten sach" in chuan_hoa(c) for c in row):
                    hdr = {}
                    for idx, c in enumerate(row):
                        t = ten_cot(c, idx)
                        if t and t not in hdr: hdr[t] = idx
                continue
            g = lambda t: row[hdr[t]] if t in hdr and hdr[t] < len(row) else None
            if not g("Tên sách") or not str(g("Tên sách")).strip():
                bo += 1; continue
            r = {t: sach_text(g(t)) for t in hdr}
            r["QĐXB"] = chuan_qdxb(g("QĐXB")); r["Lần XB"] = chuan_lan_xb(g("Lần XB"))
            rows.append(r)
        cot_la = [t for t in hdr if t not in COT_CHUAN and any(r.get(t) is not None for r in rows)]
        cols = COT_CHUAN + cot_la
        ws_o = out.create_sheet(sh)
        ws_o.append(cols)
        for r in rows: ws_o.append([r.get(c) for c in cols])
        for c in ws_o[1]: c.font = Font(bold=True)
        ws_o.freeze_panes = "C2"
        ws_o.column_dimensions["B"].width = 60; ws_o.column_dimensions["C"].width = 30
        bao_cao.append(f"{sh}: {len(rows)} sách, bỏ {bo} dòng trống, cột thêm: {cot_la or '-'}")
    out.save(ra)
    return bao_cao


if __name__ == "__main__":
    goc = sys.argv[1]
    ra = sys.argv[2] if len(sys.argv) > 2 else re.sub(r"\.xlsx$", "_chuan.xlsx", goc)
    print("\n".join(chuan_hoa_excel(goc, ra)))
    print("->", ra)
